"""
文学翻译 API 路由
支持全文翻译、四步翻译流程、RAG参考、对照编辑
"""

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, BackgroundTasks, Query, Body
from sqlalchemy import func
from sqlalchemy.orm import Session, defer, load_only, noload
from sqlalchemy.exc import OperationalError
from typing import List, Optional, AsyncGenerator
from contextvars import ContextVar
from pydantic import BaseModel
import json
import os
import io
import asyncio
import re
import tempfile
import shutil
import traceback
from app.database import get_db, SessionLocal
from app.schemas.schemas import (
    LiteraryTranslationCreate, LiteraryTranslationResponse,
    LiteraryTranslationListItem, LiteraryTranslationListPage, TranslationGroupItem,
    TranslationGroupCreate, TranslationGroupRename, TranslationGroupBulkAssign,
    LiteraryParagraphResponse, LiteraryParagraphPage,
    ParagraphUpdateRequest, LiteraryTranslationUpdate, StoryProfileUpdate,
    ReferenceDocumentCreate, ReferenceDocumentResponse,
    ReferenceDocumentListItem, ReferenceDocumentUpdate,
    WorkflowStepResponse, LiteraryTranslationWorkflowResponse,
    ExportTranslationRequest,
    ProfessionalTermCreate, ProfessionalTermUpdate, ProfessionalTermResponse,
    TranslationTermSummaryResponse
)
from app.models.models import (
    LiteraryTranslation, LiteraryParagraph, DocumentGroup,
    ReferenceDocument, LiteraryTranslationStatus,
    ProfessionalTerm, TranslationTermSummary
)
from app.core.ai_client import (
    AIClient,
    get_ai_client,
    call_for_step,
    client_for_step,
    normalize_collab_mode,
    push_collab_mode,
    reset_collab_mode,
    get_collab_mode,
)
from app.services.story_profile import (
    STORY_BATCH,
    format_digest,
    is_story_type,
    neighbor_note,
    normalize_profile,
    regenerate_story_profile,
    safe_refresh_story_profile,
    save_profile,
)
from app.services.style_agent_service import StyleAgentService

router = APIRouter(prefix="/api/literary", tags=["literary-translation"])


def build_prompt_guidance(
    translation: LiteraryTranslation,
    db: Session,
    *,
    include_story: bool = True,
) -> str:
    """翻译风格 + 故事档案，注入各步翻译 prompt。"""
    parts: List[str] = []
    style = StyleAgentService.resolve_style_block(
        db, getattr(translation, "style_agent_id", None), use_default=True
    )
    if style:
        parts.append(style)
    if include_story and is_story_type(translation.literary_type):
        story = format_digest(translation.story_profile)
        if story:
            parts.append(story)
    return "\n\n".join(parts)


def _validate_style_agent_id(db: Session, style_agent_id: Optional[int]) -> None:
    if style_agent_id is None:
        return
    if not StyleAgentService.get_agent(db, style_agent_id):
        raise HTTPException(status_code=400, detail=f"文风智能体不存在: {style_agent_id}")

class WorkflowCancelled(Exception):
    pass

# 每任务一个递增代数：终止/重开都会 +1，旧后台协程发现代数不一致即退出，避免卡在旧译文上
WORKFLOW_GENERATION: dict[int, int] = {}
_current_workflow_generation: ContextVar[Optional[int]] = ContextVar(
    "_current_workflow_generation", default=None
)


def _bump_workflow_generation(translation_id: int) -> int:
    n = WORKFLOW_GENERATION.get(translation_id, 0) + 1
    WORKFLOW_GENERATION[translation_id] = n
    return n


def _ensure_workflow_active(translation_id: int) -> None:
    """当前协程若已被终止或被新一轮启动顶替，则中断。"""
    gen = _current_workflow_generation.get()
    if gen is None:
        return
    if WORKFLOW_GENERATION.get(translation_id) != gen:
        raise WorkflowCancelled()


def _para_has_text(value: Optional[str]) -> bool:
    return bool((value or "").strip())


def _detect_resume_step(paragraphs: list) -> int:
    """根据段落已有结果决定从哪一步续跑：有译文保留，从未完成处继续。"""
    if not paragraphs:
        return 1
    if any(not _para_has_text(getattr(p, "step1_translation", None)) for p in paragraphs):
        return 1
    if any(not _para_has_text(getattr(p, "step2_verification", None)) for p in paragraphs):
        return 2
    if any(not _para_has_text(getattr(p, "step3_revision", None)) for p in paragraphs):
        return 3
    if any(not _para_has_text(getattr(p, "step4_finalization", None)) for p in paragraphs):
        return 4
    return 4

# ============================================================
# 参考文档管理
# ============================================================

@router.post("/references", response_model=ReferenceDocumentResponse)
async def create_reference_document(
    request: ReferenceDocumentCreate,
    db: Session = Depends(get_db)
):
    """创建参考文档"""
    doc = ReferenceDocument(
        name=request.name,
        file_type="txt",
        file_size=len(request.content.encode('utf-8')),
        content=request.content,
        doc_type=request.doc_type.value if hasattr(request.doc_type, 'value') else request.doc_type,
        source_lang=request.source_lang,
        target_lang=request.target_lang,
        description=request.description,
        is_active=True
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


@router.get("/references", response_model=List[ReferenceDocumentListItem])
async def list_reference_documents(
    doc_type: Optional[str] = None,
    is_active: Optional[bool] = None,
    db: Session = Depends(get_db)
):
    """获取参考文档列表"""
    query = db.query(ReferenceDocument)
    if doc_type:
        query = query.filter(ReferenceDocument.doc_type == doc_type)
    if is_active is not None:
        query = query.filter(ReferenceDocument.is_active == is_active)
    query = query.order_by(ReferenceDocument.created_at.desc())
    return query.all()


@router.get("/references/{doc_id}", response_model=ReferenceDocumentResponse)
async def get_reference_document(
    doc_id: int,
    db: Session = Depends(get_db)
):
    """获取参考文档详情"""
    doc = db.query(ReferenceDocument).filter(ReferenceDocument.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Reference document not found")
    return doc


@router.put("/references/{doc_id}", response_model=ReferenceDocumentResponse)
async def update_reference_document(
    doc_id: int,
    request: ReferenceDocumentUpdate,
    db: Session = Depends(get_db)
):
    """更新参考文档"""
    doc = db.query(ReferenceDocument).filter(ReferenceDocument.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Reference document not found")
    
    if request.name is not None:
        doc.name = request.name
    if request.content is not None:
        doc.content = request.content
        doc.file_size = len(request.content.encode('utf-8'))
    if request.doc_type is not None:
        doc.doc_type = request.doc_type.value if hasattr(request.doc_type, 'value') else request.doc_type
    if request.description is not None:
        doc.description = request.description
    if request.is_active is not None:
        doc.is_active = request.is_active
    
    db.commit()
    db.refresh(doc)
    return doc


@router.delete("/references/{doc_id}")
async def delete_reference_document(
    doc_id: int,
    db: Session = Depends(get_db)
):
    """删除参考文档"""
    doc = db.query(ReferenceDocument).filter(ReferenceDocument.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Reference document not found")
    
    db.delete(doc)
    db.commit()
    return {"message": "Reference document deleted successfully"}


@router.post("/references/upload")
async def upload_reference_document(
    file: UploadFile = File(...),
    doc_type: str = Form("general"),
    source_lang: Optional[str] = Form(None),
    target_lang: Optional[str] = Form(None),
    description: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    """上传参考文档文件，支持多种文本格式（与文学翻译上传一致）"""
    content = await file.read()
    ext = file.filename.split('.')[-1].lower() if '.' in file.filename else 'txt'
    if ext not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"不支持的文件格式: {ext}")
    try:
        text_content = parse_text_file(content, ext)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    doc = ReferenceDocument(
        name=file.filename,
        file_type=file.filename.split('.')[-1] if '.' in file.filename else 'txt',
        file_size=len(content),
        content=text_content,
        doc_type=doc_type,
        source_lang=source_lang,
        target_lang=target_lang,
        description=description,
        is_active=True
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


# ============================================================
# 文学翻译任务管理
# ============================================================

MAX_PARAGRAPH_SIZE = 2000
PARA_SEP = "\n\n"
PARA_MARKER_RE = re.compile(r"<<<PARA_(\d+)>>>")


def _split_oversized_paragraph(text: str, max_size: int = MAX_PARAGRAPH_SIZE) -> List[str]:
    """将超长段落按句子边界拆分为不超过 max_size 的块"""
    if len(text) <= max_size:
        return [text]
    sentences = re.split(r'(?<=[。！？.!?\n])', text)
    chunks, current = [], ""
    for s in sentences:
        if not s:
            continue
        if len(current) + len(s) > max_size and current:
            chunks.append(current.strip())
            current = s
        else:
            current += s
    if current.strip():
        chunks.append(current.strip())
    if not chunks:
        chunks = [text[i:i+max_size] for i in range(0, len(text), max_size)]
    return chunks


def split_text_into_paragraphs(text: str, max_size: int = MAX_PARAGRAPH_SIZE) -> List[str]:
    """将文本分割成段落，并确保每段不超过 max_size 字符"""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    # 几乎没有空行分段时：若单行大多较短，更像真实段落；硬折行长文不要按行拆
    if len(paragraphs) <= 1:
        single = [p.strip() for p in text.split("\n") if p.strip()]
        if len(single) > len(paragraphs):
            avg_len = sum(len(p) for p in single) / len(single)
            short_ratio = sum(1 for p in single if len(p) <= 120) / len(single)
            if avg_len <= 160 or short_ratio >= 0.7:
                paragraphs = single
    result = []
    for p in paragraphs:
        if len(p) > max_size:
            result.extend(_split_oversized_paragraph(p, max_size))
        else:
            result.append(p)
    return result


def _pack_marked_paragraphs(texts: List[str]) -> str:
    """用稳定标记包装段落，避免定稿后按 \\n\\n 拆分错位。"""
    parts = []
    for i, text in enumerate(texts):
        parts.append(f"<<<PARA_{i}>>>\n{(text or '').strip()}")
    return "\n".join(parts)


def _unpack_marked_paragraphs(text: str, expected: int) -> Optional[List[str]]:
    """解析带 <<<PARA_N>>> 标记的定稿结果；数量/顺序不匹配时返回 None。"""
    if not text or expected <= 0:
        return None
    matches = list(PARA_MARKER_RE.finditer(text))
    if len(matches) != expected:
        return None
    indexes = [int(m.group(1)) for m in matches]
    if indexes != list(range(expected)):
        return None
    parts: List[str] = []
    for i, match in enumerate(matches):
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        parts.append(text[start:end].strip())
    return parts


def _align_finalized_parts(
    finalized_text: str,
    fallback_texts: List[str],
) -> List[str]:
    """
    将定稿全文安全映射回各段。
    优先使用段落标记；若出现过标记但无法完整解析，直接回退 step3，避免错位。
    仅在完全没有标记时，才允许 \\n\\n 分段且数量完全一致。
    """
    expected = len(fallback_texts)
    if expected == 0:
        return []
    text = finalized_text or ""
    has_markers = bool(PARA_MARKER_RE.search(text))
    marked = _unpack_marked_paragraphs(text, expected)
    if marked is not None:
        return marked
    if has_markers:
        print(
            f"[Step4 Align] marker parse failed (expected={expected}); keep step3 revisions"
        )
        return [(t or "").strip() for t in fallback_texts]
    plain_parts = [p.strip() for p in text.split(PARA_SEP)]
    if len(plain_parts) == expected and all(plain_parts):
        return plain_parts
    print(
        f"[Step4 Align] paragraph count mismatch: got={len(plain_parts)} "
        f"expected={expected}; keep step3 revisions"
    )
    return [(t or "").strip() for t in fallback_texts]


def _sanitize_text_for_api(text: str) -> str:
    """移除空字节和控制字符，避免 Windows 上触发 OSError [Errno 22] Invalid argument"""
    if not text:
        return text
    import re
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text).strip() or text


def _safe_datetime_str(dt) -> str:
    """格式化 created_at 等时间，避免 Windows 上 strftime 触发 Errno 22"""
    if not dt:
        return ""
    try:
        return dt.strftime("%Y-%m-%d %H:%M")
    except (OSError, ValueError):
        return dt.isoformat()[:16].replace("T", " ") if hasattr(dt, "isoformat") else str(dt)


def get_reference_content(doc_ids: List[int], db: Session) -> str:
    """获取参考文档内容"""
    if not doc_ids:
        return ""
    
    docs = db.query(ReferenceDocument).filter(
        ReferenceDocument.id.in_(doc_ids),
        ReferenceDocument.is_active == True
    ).all()
    
    contents = []
    for doc in docs:
        contents.append(f"=== {doc.name} ({doc.doc_type}) ===\n{doc.content}\n")
    
    return "\n".join(contents)


def get_reference_and_requirements(translation: LiteraryTranslation, db: Session) -> str:
    """参考文档 + 用户翻译需求，供 AI 提示使用"""
    ref = get_reference_content(translation.reference_document_ids or [], db)
    if translation.user_requirements and translation.user_requirements.strip():
        req = f"\n\n## 用户翻译需求\n{translation.user_requirements.strip()}"
        ref = (ref + req) if ref else req.lstrip()
    return ref or ""


def _default_system_collab_mode() -> str:
    try:
        return normalize_collab_mode(get_collab_mode())
    except Exception:
        return "online"


def _resolve_task_collab_mode(mode: Optional[str] = None) -> str:
    if mode and str(mode).strip():
        return normalize_collab_mode(str(mode).strip())
    return _default_system_collab_mode()


@router.post("/translations", response_model=LiteraryTranslationResponse)
async def create_literary_translation(
    request: LiteraryTranslationCreate,
    db: Session = Depends(get_db)
):
    """创建文学翻译任务（大文本自动智能分段）"""
    _validate_style_agent_id(db, request.style_agent_id)
    translation = LiteraryTranslation(
        title=request.title,
        group_name=_normalize_group_name(request.group_name),
        source_text=request.source_text,
        source_lang=request.source_lang,
        target_lang=request.target_lang,
        literary_type=request.literary_type.value if hasattr(request.literary_type, 'value') else request.literary_type,
        reference_document_ids=request.reference_document_ids or [],
        user_requirements=request.user_requirements,
        style_agent_id=request.style_agent_id,
        collab_mode=_resolve_task_collab_mode(request.collab_mode),
        status=LiteraryTranslationStatus.PENDING,
        current_step=1
    )
    db.add(translation)
    _ensure_document_group(db, translation.group_name)
    db.commit()
    db.refresh(translation)
    
    paragraphs = split_text_into_paragraphs(request.source_text)
    
    for idx, para_text in enumerate(paragraphs):
        paragraph = LiteraryParagraph(
            translation_id=translation.id,
            paragraph_index=idx,
            source_text=para_text
        )
        db.add(paragraph)
    
    db.commit()
    db.refresh(translation)
    return translation


_LITERARY_CATEGORY_TYPES = ("poetry", "prose", "novel", "drama", "general")


def _normalize_group_name(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text = value.strip()
    return text or None


def _ensure_document_group(db: Session, name: Optional[str]) -> Optional[DocumentGroup]:
    """任务写入 group_name 时同步写入分组表，保证空组也可管理。"""
    normalized = _normalize_group_name(name)
    if not normalized:
        return None
    existing = db.query(DocumentGroup).filter(DocumentGroup.name == normalized).first()
    if existing:
        return existing
    group = DocumentGroup(name=normalized, sort_order=0)
    db.add(group)
    db.flush()
    return group


def _apply_translation_list_filters(
    query,
    status: Optional[str] = None,
    category: Optional[str] = None,
    group_name: Optional[str] = None,
    q: Optional[str] = None,
):
    if status:
        query = query.filter(LiteraryTranslation.status == status)
    if category == "literary":
        query = query.filter(LiteraryTranslation.literary_type.in_(_LITERARY_CATEGORY_TYPES))
    elif category == "professional":
        query = query.filter(~LiteraryTranslation.literary_type.in_(_LITERARY_CATEGORY_TYPES))
    if group_name == "__ungrouped__":
        query = query.filter(
            (LiteraryTranslation.group_name.is_(None)) | (LiteraryTranslation.group_name == "")
        )
    elif group_name:
        query = query.filter(LiteraryTranslation.group_name == group_name)
    if q and q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(LiteraryTranslation.title.ilike(like))
    return query


@router.get("/translations", response_model=LiteraryTranslationListPage)
async def list_literary_translations(
    skip: int = 0,
    limit: int = 20,
    status: Optional[str] = None,
    category: Optional[str] = Query(None, description="all|literary|professional"),
    group_name: Optional[str] = Query(
        None,
        description="分组名；传 __ungrouped__ 表示未分组",
    ),
    q: Optional[str] = Query(None, description="标题模糊搜索"),
    db: Session = Depends(get_db)
):
    """获取文学翻译任务分页列表"""
    try:
        query = db.query(LiteraryTranslation).options(
            noload(LiteraryTranslation.paragraphs),
            load_only(
                LiteraryTranslation.id,
                LiteraryTranslation.title,
                LiteraryTranslation.group_name,
                LiteraryTranslation.status,
                LiteraryTranslation.current_step,
                LiteraryTranslation.error_message,
                LiteraryTranslation.source_lang,
                LiteraryTranslation.target_lang,
                LiteraryTranslation.literary_type,
                LiteraryTranslation.collab_mode,
                LiteraryTranslation.beauty_sound_score,
                LiteraryTranslation.beauty_word_score,
                LiteraryTranslation.beauty_meaning_score,
                LiteraryTranslation.created_at,
            ),
        )
        query = _apply_translation_list_filters(
            query,
            status=status,
            category=category if category and category != "all" else None,
            group_name=group_name,
            q=q,
        )
        total = query.count()
        items = query.order_by(LiteraryTranslation.created_at.desc()).offset(skip).limit(limit).all()
        return LiteraryTranslationListPage(items=items, total=total, skip=skip, limit=limit)
    except OperationalError as e:
        err_msg = str(getattr(e, "orig", e))
        if "Unknown column" in err_msg:
            raise HTTPException(
                status_code=503,
                detail="Database schema is outdated. From project root run: mysql -u root -p aitranslator < backend/migrations/add_group_name.sql"
            )
        raise


@router.get("/translations/groups", response_model=List[TranslationGroupItem])
async def list_translation_groups(
    category: Optional[str] = Query(None, description="all|literary|professional"),
    db: Session = Depends(get_db),
):
    """获取自定义分组列表（含文档数量；空组也会返回）"""
    count_query = db.query(
        LiteraryTranslation.group_name,
        func.count(LiteraryTranslation.id),
    )
    if category == "literary":
        count_query = count_query.filter(LiteraryTranslation.literary_type.in_(_LITERARY_CATEGORY_TYPES))
    elif category == "professional":
        count_query = count_query.filter(~LiteraryTranslation.literary_type.in_(_LITERARY_CATEGORY_TYPES))
    count_map = {
        ((name or "").strip() or ""): count
        for name, count in count_query.group_by(LiteraryTranslation.group_name).all()
    }

    # 同步已有任务分组名进分组表（兼容升级前数据）
    for name in list(count_map.keys()):
        if name:
            _ensure_document_group(db, name)
    db.commit()

    groups = db.query(DocumentGroup).order_by(DocumentGroup.sort_order.asc(), DocumentGroup.name.asc()).all()
    items: List[TranslationGroupItem] = [
        TranslationGroupItem(
            id=g.id,
            name=g.name,
            count=count_map.get(g.name, 0),
            sort_order=g.sort_order or 0,
        )
        for g in groups
    ]
    ungrouped = count_map.get("", 0)
    if ungrouped:
        items.append(TranslationGroupItem(id=None, name="", count=ungrouped, sort_order=10_000))
    return items


@router.post("/translations/groups", response_model=TranslationGroupItem)
async def create_translation_group(
    request: TranslationGroupCreate,
    db: Session = Depends(get_db),
):
    """新建自定义分组（可为空组）"""
    name = _normalize_group_name(request.name)
    if not name:
        raise HTTPException(status_code=400, detail="分组名称不能为空")
    exists = db.query(DocumentGroup).filter(DocumentGroup.name == name).first()
    if exists:
        raise HTTPException(status_code=400, detail="分组已存在")
    group = DocumentGroup(name=name, sort_order=0)
    db.add(group)
    db.commit()
    db.refresh(group)
    return TranslationGroupItem(id=group.id, name=group.name, count=0, sort_order=group.sort_order or 0)


@router.put("/translations/groups/{group_id}", response_model=TranslationGroupItem)
async def rename_translation_group(
    group_id: int,
    request: TranslationGroupRename,
    db: Session = Depends(get_db),
):
    """重命名分组，并同步更新该组下所有文档"""
    group = db.query(DocumentGroup).filter(DocumentGroup.id == group_id).first()
    if not group:
        raise HTTPException(status_code=404, detail="分组不存在")
    new_name = _normalize_group_name(request.name)
    if not new_name:
        raise HTTPException(status_code=400, detail="分组名称不能为空")
    if new_name != group.name:
        conflict = db.query(DocumentGroup).filter(
            DocumentGroup.name == new_name,
            DocumentGroup.id != group_id,
        ).first()
        if conflict:
            raise HTTPException(status_code=400, detail="目标分组名已存在")
        old_name = group.name
        group.name = new_name
        db.query(LiteraryTranslation).filter(LiteraryTranslation.group_name == old_name).update(
            {LiteraryTranslation.group_name: new_name},
            synchronize_session=False,
        )
        db.commit()
        db.refresh(group)
    count = db.query(func.count(LiteraryTranslation.id)).filter(
        LiteraryTranslation.group_name == group.name
    ).scalar() or 0
    return TranslationGroupItem(id=group.id, name=group.name, count=count, sort_order=group.sort_order or 0)


@router.delete("/translations/groups/{group_id}")
async def delete_translation_group(
    group_id: int,
    clear_docs: bool = Query(True, description="是否把该组文档改为未分组"),
    db: Session = Depends(get_db),
):
    """删除自定义分组"""
    group = db.query(DocumentGroup).filter(DocumentGroup.id == group_id).first()
    if not group:
        raise HTTPException(status_code=404, detail="分组不存在")
    affected = 0
    if clear_docs:
        affected = db.query(LiteraryTranslation).filter(
            LiteraryTranslation.group_name == group.name
        ).update({LiteraryTranslation.group_name: None}, synchronize_session=False)
    db.delete(group)
    db.commit()
    return {"message": "分组已删除", "cleared_docs": affected}


@router.post("/translations/groups/bulk-assign")
async def bulk_assign_translation_group(
    request: TranslationGroupBulkAssign,
    db: Session = Depends(get_db),
):
    """批量设置文档分组"""
    group_name = _normalize_group_name(request.group_name)
    if group_name:
        _ensure_document_group(db, group_name)
    updated = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id.in_(request.translation_ids)
    ).update({LiteraryTranslation.group_name: group_name}, synchronize_session=False)
    db.commit()
    return {"message": "已更新分组", "updated": updated, "group_name": group_name}


_HEAVY_TEXT_COLUMNS = (
    LiteraryTranslation.source_text,
    LiteraryTranslation.step1_translation,
    LiteraryTranslation.step2_verification,
    LiteraryTranslation.step3_revision,
    LiteraryTranslation.step4_finalization,
    LiteraryTranslation.final_translation,
)


@router.get("/translations/{translation_id}", response_model=LiteraryTranslationResponse)
async def get_literary_translation(
    translation_id: int,
    include_paragraphs: bool = True,
    include_source: bool = True,
    db: Session = Depends(get_db)
):
    """获取文学翻译任务详情。浏览长文时传 include_paragraphs=false&include_source=false，再分页拉段落。"""
    options = []
    if not include_paragraphs:
        options.append(noload(LiteraryTranslation.paragraphs))
    if not include_source:
        options.extend(defer(column) for column in _HEAVY_TEXT_COLUMNS)

    translation = db.query(LiteraryTranslation).options(*options).filter(
        LiteraryTranslation.id == translation_id
    ).first()

    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")

    paragraph_total = db.query(func.count(LiteraryParagraph.id)).filter(
        LiteraryParagraph.translation_id == translation_id
    ).scalar() or 0

    if include_source:
        flags = {
            "has_step2": bool(translation.step2_verification),
            "has_step3": bool(translation.step3_revision),
            "has_step4": bool(translation.step4_finalization),
        }
    else:
        row = db.query(
            LiteraryTranslation.step2_verification.isnot(None),
            LiteraryTranslation.step3_revision.isnot(None),
            LiteraryTranslation.step4_finalization.isnot(None),
        ).filter(LiteraryTranslation.id == translation_id).one()
        flags = {"has_step2": bool(row[0]), "has_step3": bool(row[1]), "has_step4": bool(row[2])}

    return LiteraryTranslationResponse(
        id=translation.id,
        title=translation.title,
        group_name=getattr(translation, "group_name", None),
        source_text=translation.source_text if include_source else "",
        step1_translation=translation.step1_translation if include_source else None,
        step2_verification=translation.step2_verification if include_source else None,
        step3_revision=translation.step3_revision if include_source else None,
        step4_finalization=translation.step4_finalization if include_source else None,
        final_translation=translation.final_translation if include_source else None,
        current_step=translation.current_step or 1,
        status=translation.status,
        source_lang=translation.source_lang,
        target_lang=translation.target_lang,
        literary_type=translation.literary_type,
        collab_mode=getattr(translation, "collab_mode", None) or "online",
        beauty_sound_score=translation.beauty_sound_score,
        beauty_word_score=translation.beauty_word_score,
        beauty_meaning_score=translation.beauty_meaning_score,
        ai_provider=translation.ai_provider,
        ai_model=translation.ai_model,
        reference_document_ids=translation.reference_document_ids or [],
        user_requirements=translation.user_requirements,
        style_agent_id=getattr(translation, "style_agent_id", None),
        created_at=translation.created_at,
        updated_at=translation.updated_at,
        completed_at=translation.completed_at,
        paragraphs=list(translation.paragraphs) if include_paragraphs else None,
        error_message=translation.error_message,
        has_step2=flags["has_step2"],
        has_step3=flags["has_step3"],
        has_step4=flags["has_step4"],
        paragraph_total=paragraph_total,
    )


@router.put("/translations/{translation_id}", response_model=LiteraryTranslationResponse)
async def update_literary_translation(
    translation_id: int,
    request: LiteraryTranslationUpdate,
    db: Session = Depends(get_db)
):
    """更新文学翻译任务"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    
    if request.title is not None:
        translation.title = request.title
    fields_set = getattr(request, "model_fields_set", None) or set()
    if "group_name" in fields_set:
        translation.group_name = _normalize_group_name(request.group_name)
        _ensure_document_group(db, translation.group_name)
    if request.source_text is not None:
        translation.source_text = request.source_text
    if request.final_translation is not None:
        translation.final_translation = request.final_translation
    if request.status is not None:
        allowed = ("pending", "translating", "verifying", "revising", "finalizing", "completed", "failed")
        if request.status not in allowed:
            raise HTTPException(status_code=400, detail=f"status must be one of: {allowed}")
        translation.status = request.status
    if "style_agent_id" in fields_set:
        _validate_style_agent_id(db, request.style_agent_id)
        translation.style_agent_id = request.style_agent_id
    if request.user_requirements is not None:
        translation.user_requirements = request.user_requirements

    db.commit()
    db.refresh(translation)
    return translation


@router.delete("/translations/{translation_id}")
async def delete_literary_translation(
    translation_id: int,
    db: Session = Depends(get_db)
):
    """删除文学翻译任务"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    
    db.delete(translation)
    db.commit()
    return {"message": "Translation deleted successfully"}


# ============================================================
# 四步翻译流程 API（内部执行函数供 run_all 复用）
# ============================================================

async def _execute_step1(translation_id: int, db: Session) -> None:
    """执行第一步：初译（逐段翻译，每段完成即保存）"""
    translation = db.query(LiteraryTranslation).filter(LiteraryTranslation.id == translation_id).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    translation.status = LiteraryTranslationStatus.TRANSLATING
    db.commit()
    client, _role = client_for_step(1)
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    reference_content = get_reference_and_requirements(translation, db)
    story_mode = is_story_type(translation.literary_type)
    if story_mode and paragraphs:
        await safe_refresh_story_profile(client, translation, paragraphs[:STORY_BATCH], db, include_translation=False)
        db.refresh(translation)
    full_step1 = []
    pending_story = []
    for index, para in enumerate(paragraphs):
        _ensure_workflow_active(translation_id)
        # 重启续跑：已有初译的段落直接保留
        if _para_has_text(para.step1_translation):
            full_step1.append(para.step1_translation)
            if not _para_has_text(para.translated_text):
                para.translated_text = para.step1_translation
                db.commit()
            continue
        source_text = _sanitize_text_for_api(para.source_text or "")
        ref_safe = _sanitize_text_for_api(reference_content) if reference_content else ""
        guidance = build_prompt_guidance(translation, db, include_story=story_mode)
        neighbor = neighbor_note(paragraphs, index) if story_mode else ""
        result = await call_for_step(
            1,
            lambda c, text=source_text, ref=ref_safe, guide=guidance, note=neighbor: c.literary_translate(
                text, translation.source_lang, translation.target_lang,
                translation.literary_type, ref,
                guidance=guide or None,
                neighbor_context=note or None,
            ),
        )
        cleaned = AIClient._strip_translation_wrappers(result)
        para.step1_translation = cleaned
        para.translated_text = cleaned
        full_step1.append(cleaned)
        db.commit()
        if story_mode:
            pending_story.append(para)
            if len(pending_story) >= STORY_BATCH:
                await safe_refresh_story_profile(client, translation, pending_story, db, include_translation=True)
                db.refresh(translation)
                pending_story = []
    if story_mode and pending_story:
        await safe_refresh_story_profile(client, translation, pending_story, db, include_translation=True)
        db.refresh(translation)
    translation.step1_translation = "\n\n".join(full_step1)
    translation.current_step = 2
    translation.status = LiteraryTranslationStatus.VERIFYING
    db.commit()


async def _execute_step2(translation_id: int, db: Session) -> None:
    """执行第二步：校验"""
    translation = db.query(LiteraryTranslation).filter(LiteraryTranslation.id == translation_id).first()
    if not translation or translation.current_step < 2:
        raise HTTPException(status_code=400, detail="Please complete step 1 first")
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    full_step2 = []
    total_sound, total_word, total_meaning = 0, 0, 0
    guidance = build_prompt_guidance(translation, db, include_story=True)
    for para in paragraphs:
        _ensure_workflow_active(translation_id)
        if _para_has_text(para.step2_verification):
            verified = para.step2_verification
            full_step2.append(verified)
            total_sound += para.beauty_sound_score or 7.0
            total_word += para.beauty_word_score or 7.0
            total_meaning += para.beauty_meaning_score or 7.0
            if not _para_has_text(para.translated_text):
                para.translated_text = verified
                db.commit()
            continue
        result = await call_for_step(
            2,
            lambda c, p=para, guide=guidance: c.literary_verify(
                p.source_text, p.step1_translation or p.translated_text or "",
                translation.source_lang, translation.target_lang, translation.literary_type,
                guidance=guide or None,
            ),
        )
        verified = AIClient._strip_translation_wrappers(
            result.get("verified_translation", para.step1_translation) or ""
        )
        para.step2_verification = verified
        para.translated_text = verified
        para.beauty_sound_score = result.get("beauty_sound_score", 7.0)
        para.beauty_word_score = result.get("beauty_word_score", 7.0)
        para.beauty_meaning_score = result.get("beauty_meaning_score", 7.0)
        total_sound += para.beauty_sound_score
        total_word += para.beauty_word_score
        total_meaning += para.beauty_meaning_score
        full_step2.append(verified)
        db.commit()
    translation.step2_verification = "\n\n".join(full_step2)
    translation.current_step = 3
    translation.status = LiteraryTranslationStatus.REVISING
    if paragraphs:
        translation.beauty_sound_score = total_sound / len(paragraphs)
        translation.beauty_word_score = total_word / len(paragraphs)
        translation.beauty_meaning_score = total_meaning / len(paragraphs)
    db.commit()


async def _execute_step3(translation_id: int, db: Session) -> None:
    """执行第三步：修改"""
    translation = db.query(LiteraryTranslation).filter(LiteraryTranslation.id == translation_id).first()
    if not translation or translation.current_step < 3:
        raise HTTPException(status_code=400, detail="Please complete step 2 first")
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    full_step3 = []
    guidance = build_prompt_guidance(translation, db, include_story=True)
    for para in paragraphs:
        _ensure_workflow_active(translation_id)
        if _para_has_text(para.step3_revision):
            full_step3.append(para.step3_revision)
            if not _para_has_text(para.translated_text):
                para.translated_text = para.step3_revision
                db.commit()
            continue
        verification_analysis = {
            "issues_found": [], "suggestions": [],
            "beauty_sound_score": para.beauty_sound_score,
            "beauty_word_score": para.beauty_word_score,
            "beauty_meaning_score": para.beauty_meaning_score,
        }
        result = await call_for_step(
            3,
            lambda c, p=para, analysis=verification_analysis, guide=guidance: c.literary_revise(
                p.source_text, p.step2_verification or p.translated_text or "",
                analysis, translation.source_lang, translation.target_lang, translation.literary_type,
                guidance=guide or None,
            ),
        )
        revised = AIClient._strip_translation_wrappers(
            result.get("revised_translation", para.step2_verification) or ""
        )
        para.step3_revision = revised
        para.translated_text = revised
        full_step3.append(revised)
        db.commit()
    translation.step3_revision = "\n\n".join(full_step3)
    translation.current_step = 4
    translation.status = LiteraryTranslationStatus.FINALIZING
    db.commit()


STEP4_BATCH_CHARS = 6000


def _apply_finalized_parts(paragraphs: list, finalized_parts: List[str]) -> None:
    for para, text in zip(paragraphs, finalized_parts):
        cleaned = AIClient._strip_translation_wrappers(text or "")
        para.step4_finalization = cleaned
        para.translated_text = cleaned


def _split_paragraphs_into_batches(paragraphs: list, size_fn, max_chars: int = STEP4_BATCH_CHARS) -> list:
    """按字符量把段落切成批次。"""
    batches: list[list] = []
    current_batch: list = []
    current_size = 0
    for para in paragraphs:
        p_size = size_fn(para)
        if current_size + p_size > max_chars and current_batch:
            batches.append(current_batch)
            current_batch = [para]
            current_size = p_size
        else:
            current_batch.append(para)
            current_size += p_size
    if current_batch:
        batches.append(current_batch)
    return batches


async def _finalize_paragraph_batch(
    translation,
    batch: list,
    db: Session,
) -> None:
    """对一批段落做定稿，并用标记对齐写回，避免原文/译文错位。"""
    fallback_texts = [
        (p.step3_revision or p.translated_text or "") for p in batch
    ]
    batch_source = _pack_marked_paragraphs([p.source_text or "" for p in batch])
    batch_revised = _pack_marked_paragraphs(fallback_texts)
    guidance = build_prompt_guidance(translation, db, include_story=True) or None
    result = await call_for_step(
        4,
        lambda c, src=batch_source, revised=batch_revised, guide=guidance, count=len(batch): c.literary_finalize(
            src,
            revised,
            translation.source_lang,
            translation.target_lang,
            translation.literary_type,
            guidance=guide,
            paragraph_count=count,
        ),
    )
    finalized_text = result.get("final_translation", batch_revised)
    finalized_parts = _align_finalized_parts(finalized_text, fallback_texts)
    _apply_finalized_parts(batch, finalized_parts)
    db.commit()


async def _holistic_errata_batch(
    translation,
    batch: list,
    db: Session,
) -> None:
    """对一批已定稿段落做整体勘误（常识/文化/习俗用语等）。"""
    fallback_texts = [
        (p.step4_finalization or p.step3_revision or p.translated_text or "") for p in batch
    ]
    batch_source = _pack_marked_paragraphs([p.source_text or "" for p in batch])
    batch_draft = _pack_marked_paragraphs(fallback_texts)
    guidance = build_prompt_guidance(translation, db, include_story=True) or None
    result = await call_for_step(
        4,
        lambda c, src=batch_source, draft=batch_draft, guide=guidance, count=len(batch): c.literary_holistic_errata(
            src,
            draft,
            translation.source_lang,
            translation.target_lang,
            translation.literary_type,
            guidance=guide,
            paragraph_count=count,
        ),
    )
    corrected = result.get("corrected_translation", batch_draft)
    corrected_parts = _align_finalized_parts(corrected, fallback_texts)
    _apply_finalized_parts(batch, corrected_parts)
    errata = result.get("errata") or []
    summary = (result.get("summary") or "").strip()
    if errata or summary:
        print(
            f"[HolisticErrata] translation={translation.id} paras="
            f"{batch[0].paragraph_index}-{batch[-1].paragraph_index} "
            f"fixes={len(errata)} summary={summary[:120]}"
        )
    db.commit()


async def _execute_step4(translation_id: int, db: Session) -> None:
    """执行第四步：定稿 + 整体勘误（常识/文化/习俗用语等）"""
    translation = db.query(LiteraryTranslation).filter(LiteraryTranslation.id == translation_id).first()
    if not translation or translation.current_step < 4:
        raise HTTPException(status_code=400, detail="Please complete step 3 first")
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()

    # 1) 定稿：只处理尚未完成的段落
    pending = [p for p in paragraphs if not _para_has_text(p.step4_finalization)]
    if pending:
        total_chars = sum(len(p.step3_revision or p.translated_text or "") for p in pending)
        if total_chars <= STEP4_BATCH_CHARS:
            _ensure_workflow_active(translation_id)
            await _finalize_paragraph_batch(translation, pending, db)
        else:
            for batch in _split_paragraphs_into_batches(
                pending,
                lambda p: len(p.step3_revision or p.translated_text or ""),
            ):
                _ensure_workflow_active(translation_id)
                await _finalize_paragraph_batch(translation, batch, db)

    # 2) 整体勘误：以全文/分批视角订正常识、文化、习俗用语等
    # 刷新段落，确保使用最新定稿文本
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    if paragraphs:
        total_chars = sum(
            len(p.step4_finalization or p.translated_text or "") for p in paragraphs
        )
        if total_chars <= STEP4_BATCH_CHARS:
            _ensure_workflow_active(translation_id)
            await _holistic_errata_batch(translation, paragraphs, db)
        else:
            for batch in _split_paragraphs_into_batches(
                paragraphs,
                lambda p: len(p.step4_finalization or p.translated_text or ""),
            ):
                _ensure_workflow_active(translation_id)
                await _holistic_errata_batch(translation, batch, db)

    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    finalized_all = PARA_SEP.join(
        (p.step4_finalization or p.translated_text or "") for p in paragraphs
    )
    translation.step4_finalization = finalized_all
    translation.final_translation = finalized_all
    translation.current_step = 4
    translation.status = LiteraryTranslationStatus.COMPLETED
    from datetime import datetime
    translation.completed_at = datetime.now()
    db.commit()
    try:
        await auto_extract_terms_after_finalize(translation_id, db)
    except Exception as e:
        print(f"[Auto Extract Terms] Error: {e}")


def _format_error_debug(exc: BaseException, step: int, step_name: str) -> str:
    """将异常格式化为带调试信息的 error_message（便于排查 [Errno 22] 等）"""
    tb_lines = traceback.format_exception(type(exc), exc, exc.__traceback__)
    tb_str = "".join(tb_lines).strip()
    # 限制长度，避免 DB 字段过长；保留最后约 2500 字符（通常含关键堆栈）
    max_tb = 2500
    if len(tb_str) > max_tb:
        tb_str = "...\n" + tb_str[-max_tb:]
    return (
        f"[Step {step} - {step_name}] {type(exc).__name__}: {exc}\n\n"
        f"Traceback:\n{tb_str}"
    )


async def _run_workflow_background(translation_id: int, generation: int):
    """后台依次执行四步翻译流程：初译 → 校验 → 修改 → 定稿"""
    db = SessionLocal()
    steps = [
        (1, "初译", _execute_step1),
        (2, "校验", _execute_step2),
        (3, "修改", _execute_step3),
        (4, "定稿", _execute_step4),
    ]
    mode_token = None
    gen_token = _current_workflow_generation.set(generation)
    try:
        # 已被更新的启动顶替则直接退出，避免写回旧结果
        if WORKFLOW_GENERATION.get(translation_id) != generation:
            return
        t0 = db.query(LiteraryTranslation).filter(LiteraryTranslation.id == translation_id).first()
        mode_token = push_collab_mode(getattr(t0, "collab_mode", None) if t0 else None)
        start_step = (getattr(t0, "current_step", 1) or 1) if t0 else 1
        for step_num, step_name, step_fn in steps:
            if step_num < start_step:
                continue
            _ensure_workflow_active(translation_id)
            await step_fn(translation_id, db)
    except WorkflowCancelled:
        # 仅当前这一轮仍有效时才标失败；被重开顶替时不要覆盖新任务状态
        if WORKFLOW_GENERATION.get(translation_id) != generation:
            return
        try:
            translation = db.query(LiteraryTranslation).filter(
                LiteraryTranslation.id == translation_id
            ).first()
            if translation:
                translation.status = LiteraryTranslationStatus.FAILED
                if not translation.error_message:
                    translation.error_message = "Cancelled by user"
                db.commit()
        except Exception:
            pass
    except Exception as e:
        if WORKFLOW_GENERATION.get(translation_id) != generation:
            return
        # 确定失败步骤（当前步骤尚未完成）
        failed_step = 1
        failed_name = "初译"
        try:
            t = db.query(LiteraryTranslation).filter(
                LiteraryTranslation.id == translation_id
            ).first()
            if t is not None:
                # current_step 是“正在做”的步骤，失败时就是该步
                failed_step = getattr(t, "current_step", 1) or 1
                for sn, sname, _ in steps:
                    if sn == failed_step:
                        failed_name = sname
                        break
        except Exception:
            pass
        try:
            translation = db.query(LiteraryTranslation).filter(
                LiteraryTranslation.id == translation_id
            ).first()
            if translation:
                translation.status = LiteraryTranslationStatus.FAILED
                translation.error_message = _format_error_debug(e, failed_step, failed_name)
                db.commit()
        except Exception:
            pass
        traceback.print_exc()
        print(f"[Workflow] Translation {translation_id} failed: {e}")
    finally:
        _current_workflow_generation.reset(gen_token)
        if mode_token is not None:
            reset_collab_mode(mode_token)
        db.close()


class WorkflowStartRequest(BaseModel):
    collab_mode: Optional[str] = None


@router.post("/translations/{translation_id}/workflow/start")
async def start_translation_workflow(
    translation_id: int,
    request: WorkflowStartRequest = Body(default_factory=WorkflowStartRequest),
    db: Session = Depends(get_db)
):
    """启动四步翻译流程（后台执行：初译 → 校验 → 修改 → 定稿），立即返回，前端轮询状态"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")

    # 允许失败/完成/卡住任务重开；bump 代数会顶替仍在跑的旧后台协程
    generation = _bump_workflow_generation(translation_id)

    mode = request.collab_mode if request else None
    translation.collab_mode = _resolve_task_collab_mode(mode or getattr(translation, "collab_mode", None))

    # 保留段落已有译文，从未完成的步骤/段落继续
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    resume_step = _detect_resume_step(paragraphs)
    status_by_step = {
        1: LiteraryTranslationStatus.TRANSLATING,
        2: LiteraryTranslationStatus.VERIFYING,
        3: LiteraryTranslationStatus.REVISING,
        4: LiteraryTranslationStatus.FINALIZING,
    }
    translation.current_step = resume_step
    translation.status = status_by_step.get(resume_step, LiteraryTranslationStatus.TRANSLATING)
    # 任务级汇总按已完成步骤回填，便于进度展示；段落级内容全部保留
    translation.step1_translation = (
        "\n\n".join(p.step1_translation or "" for p in paragraphs if _para_has_text(p.step1_translation))
        if resume_step > 1 else None
    )
    translation.step2_verification = (
        "\n\n".join(p.step2_verification or "" for p in paragraphs if _para_has_text(p.step2_verification))
        if resume_step > 2 else None
    )
    translation.step3_revision = (
        "\n\n".join(p.step3_revision or "" for p in paragraphs if _para_has_text(p.step3_revision))
        if resume_step > 3 else None
    )
    translation.step4_finalization = None
    translation.final_translation = None
    translation.completed_at = None
    translation.error_message = None

    db.commit()

    asyncio.create_task(_run_workflow_background(translation_id, generation))

    return {
        "message": "翻译流程已启动",
        "status": translation.status,
        "collab_mode": translation.collab_mode,
        "resume_step": resume_step,
    }

@router.post("/translations/{translation_id}/workflow/stop")
async def stop_translation_workflow(
    translation_id: int,
    db: Session = Depends(get_db)
):
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    running_statuses = {
        LiteraryTranslationStatus.TRANSLATING,
        LiteraryTranslationStatus.VERIFYING,
        LiteraryTranslationStatus.REVISING,
        LiteraryTranslationStatus.FINALIZING,
    }
    if translation.status not in running_statuses:
        return {"message": "当前未在运行", "status": translation.status}
    # bump 代数，令后台协程在下一段/下一步时自行退出
    _bump_workflow_generation(translation_id)
    translation.status = LiteraryTranslationStatus.FAILED
    translation.error_message = "Cancelled by user"
    db.commit()
    return {"message": "已终止", "status": "failed"}

async def _run_batch_workflow_background(translation_ids: list):
    """依次执行多个任务的翻译流程，前一个完成后才开始下一个"""
    for tid in translation_ids:
        db = SessionLocal()
        try:
            translation = db.query(LiteraryTranslation).filter(
                LiteraryTranslation.id == tid
            ).first()
            if not translation:
                continue
            generation = _bump_workflow_generation(tid)
            paragraphs = db.query(LiteraryParagraph).filter(
                LiteraryParagraph.translation_id == tid
            ).order_by(LiteraryParagraph.paragraph_index).all()
            resume_step = _detect_resume_step(paragraphs)
            status_by_step = {
                1: LiteraryTranslationStatus.TRANSLATING,
                2: LiteraryTranslationStatus.VERIFYING,
                3: LiteraryTranslationStatus.REVISING,
                4: LiteraryTranslationStatus.FINALIZING,
            }
            translation.current_step = resume_step
            translation.status = status_by_step.get(resume_step, LiteraryTranslationStatus.TRANSLATING)
            translation.step1_translation = None
            translation.step2_verification = None
            translation.step3_revision = None
            translation.step4_finalization = None
            translation.final_translation = None
            translation.completed_at = None
            translation.error_message = None
            db.commit()
        except Exception:
            db.close()
            continue
        finally:
            db.close()

        try:
            await _run_workflow_background(tid, generation)
        except Exception as e:
            print(f"[BatchWorkflow] Translation {tid} failed: {e}")


@router.post("/translations/batch/workflow/start")
async def start_batch_workflow(
    request: dict,
    db: Session = Depends(get_db)
):
    """批量启动翻译流程，按顺序依次处理多个任务"""
    translation_ids = request.get("translation_ids", [])
    if not translation_ids:
        raise HTTPException(status_code=400, detail="请选择至少一个任务")
    batch_mode = request.get("collab_mode")
    resolved_mode = _resolve_task_collab_mode(batch_mode) if batch_mode else None

    running_statuses = {
        LiteraryTranslationStatus.TRANSLATING,
        LiteraryTranslationStatus.VERIFYING,
        LiteraryTranslationStatus.REVISING,
        LiteraryTranslationStatus.FINALIZING,
    }

    valid_ids = []
    for tid in translation_ids:
        t = db.query(LiteraryTranslation).filter(LiteraryTranslation.id == tid).first()
        if not t:
            continue
        if t.status in running_statuses:
            continue
        if resolved_mode:
            t.collab_mode = resolved_mode
        elif not getattr(t, "collab_mode", None):
            t.collab_mode = _default_system_collab_mode()
        t.status = LiteraryTranslationStatus.PENDING
        valid_ids.append(tid)

    if not valid_ids:
        raise HTTPException(status_code=400, detail="没有可启动的任务")

    db.commit()
    asyncio.create_task(_run_batch_workflow_background(valid_ids))

    return {"message": f"已加入队列 {len(valid_ids)} 个任务", "translation_ids": valid_ids}


@router.get("/translations/{translation_id}/story")
async def get_story_profile(
    translation_id: int,
    db: Session = Depends(get_db)
):
    """获取故事结构档案"""
    translation = db.query(LiteraryTranslation).options(
        noload(LiteraryTranslation.paragraphs),
        *[defer(column) for column in _HEAVY_TEXT_COLUMNS],
    ).filter(LiteraryTranslation.id == translation_id).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    return {
        "translation_id": translation.id,
        "title": translation.title,
        "literary_type": translation.literary_type,
        "status": translation.status,
        "applicable": is_story_type(translation.literary_type),
        "profile": normalize_profile(translation.story_profile),
    }


@router.put("/translations/{translation_id}/story")
async def update_story_profile_endpoint(
    translation_id: int,
    request: StoryProfileUpdate,
    db: Session = Depends(get_db)
):
    """保存用户修订后的故事结构档案。已确认译名会写入词库。"""
    translation = db.query(LiteraryTranslation).options(
        noload(LiteraryTranslation.paragraphs),
        *[defer(column) for column in _HEAVY_TEXT_COLUMNS],
    ).filter(LiteraryTranslation.id == translation_id).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    if not is_story_type(translation.literary_type):
        raise HTTPException(status_code=400, detail="当前文本类型不整理故事结构")
    profile = save_profile(translation, request.profile, db)
    return {"profile": profile}


@router.post("/translations/{translation_id}/story/regenerate")
async def regenerate_story_profile_endpoint(
    translation_id: int,
    db: Session = Depends(get_db),
):
    """根据全部段落原文重新生成故事结构（保留已确认译名）。"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    if not is_story_type(translation.literary_type):
        raise HTTPException(status_code=400, detail="当前文本类型不整理故事结构")

    client = get_ai_client()
    try:
        profile = await regenerate_story_profile(
            client, translation, db, include_translation=False
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e)) from e
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"重新生成失败: {e}") from e

    return {
        "translation_id": translation.id,
        "title": translation.title,
        "literary_type": translation.literary_type,
        "status": translation.status,
        "applicable": True,
        "profile": profile,
    }


@router.get("/translations/{translation_id}/workflow", response_model=LiteraryTranslationWorkflowResponse)
async def get_workflow_status(
    translation_id: int,
    db: Session = Depends(get_db)
):
    """获取工作流状态（供前端轮询）"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()

    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")

    status_to_running_step = {
        "translating": 1, "verifying": 2, "revising": 3, "finalizing": 4,
    }
    running_step = status_to_running_step.get(translation.status, 0)

    step_fields = [
        (1, "翻译", translation.step1_translation),
        (2, "校验", translation.step2_verification),
        (3, "修改", translation.step3_revision),
        (4, "定稿勘误", translation.step4_finalization),
    ]

    steps = []
    for step_num, step_name, result in step_fields:
        if result:
            step_status = "completed"
        elif step_num == running_step:
            step_status = "processing"
        else:
            step_status = "pending"
        truncated = result[:200] + "..." if result and len(result) > 200 else result
        steps.append(WorkflowStepResponse(
            step=step_num, step_name=step_name, status=step_status, result=truncated,
        ))

    para_total = 0
    para_done = 0
    if running_step > 0:
        all_paras = db.query(LiteraryParagraph).filter(
            LiteraryParagraph.translation_id == translation_id
        ).all()
        para_total = len(all_paras)
        step_field_map = {1: "step1_translation", 2: "step2_verification", 3: "step3_revision", 4: "step4_finalization"}
        field = step_field_map.get(running_step)
        if field:
            para_done = sum(1 for p in all_paras if getattr(p, field, None))

    return LiteraryTranslationWorkflowResponse(
        translation_id=translation_id,
        current_step=translation.current_step,
        overall_status=translation.status,
        steps=steps,
        paragraph_total=para_total,
        paragraph_done=para_done,
    )


# ============================================================
# 段落管理
# ============================================================

@router.get("/translations/{translation_id}/paragraphs", response_model=LiteraryParagraphPage)
async def get_paragraphs(
    translation_id: int,
    skip: int = Query(0, ge=0),
    limit: int = Query(30, ge=1, le=200),
    db: Session = Depends(get_db)
):
    """分页获取段落，避免长文本一次加载全部正文"""
    exists = db.query(LiteraryTranslation.id).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    if not exists:
        raise HTTPException(status_code=404, detail="Translation not found")

    total = db.query(func.count(LiteraryParagraph.id)).filter(
        LiteraryParagraph.translation_id == translation_id
    ).scalar() or 0
    items = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).offset(skip).limit(limit).all()

    return LiteraryParagraphPage(items=items, total=total, skip=skip, limit=limit)


@router.put("/paragraphs/{paragraph_id}", response_model=LiteraryParagraphResponse)
async def update_paragraph(
    paragraph_id: int,
    request: ParagraphUpdateRequest,
    db: Session = Depends(get_db)
):
    """更新段落译文（用户编辑）"""
    paragraph = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.id == paragraph_id
    ).first()
    
    if not paragraph:
        raise HTTPException(status_code=404, detail="Paragraph not found")
    
    paragraph.user_edited_text = request.user_edited_text
    paragraph.translated_text = request.user_edited_text
    paragraph.is_edited = True
    
    db.commit()
    db.refresh(paragraph)
    
    # 更新全文
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == paragraph.translation_id
    ).first()
    
    if translation:
        all_paras = db.query(LiteraryParagraph).filter(
            LiteraryParagraph.translation_id == translation.id
        ).order_by(LiteraryParagraph.paragraph_index).all()
        
        translation.final_translation = "\n\n".join([
            p.user_edited_text or p.translated_text or "" for p in all_paras
        ])
        db.commit()
    
    return paragraph


@router.post("/paragraphs/{paragraph_id}/retranslate", response_model=LiteraryParagraphResponse)
async def retranslate_paragraph(
    paragraph_id: int,
    db: Session = Depends(get_db),
):
    """对单个段落重新跑完整四步翻译（初译→校验→润色→定稿）"""
    paragraph = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.id == paragraph_id
    ).first()
    if not paragraph:
        raise HTTPException(status_code=404, detail="Paragraph not found")

    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == paragraph.translation_id
    ).first()
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")

    source_text = _sanitize_text_for_api(paragraph.source_text or "")
    if not source_text.strip():
        raise HTTPException(status_code=400, detail="该段原文为空，无法重译")

    all_paras = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation.id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    index = next((i for i, p in enumerate(all_paras) if p.id == paragraph.id), 0)

    reference_content = get_reference_and_requirements(translation, db)
    ref_safe = _sanitize_text_for_api(reference_content) if reference_content else ""
    story_mode = is_story_type(translation.literary_type)
    guidance = build_prompt_guidance(translation, db, include_story=story_mode)
    neighbor = neighbor_note(all_paras, index) if story_mode else ""
    mode_token = push_collab_mode(getattr(translation, "collab_mode", None))

    try:
        step1 = await call_for_step(
            1,
            lambda c: c.literary_translate(
                source_text, translation.source_lang, translation.target_lang,
                translation.literary_type, ref_safe or None,
                guidance=guidance or None,
                neighbor_context=neighbor or None,
            ),
        )
        step2_result = await call_for_step(
            2,
            lambda c: c.literary_verify(
                source_text, step1, translation.source_lang, translation.target_lang,
                translation.literary_type, guidance=guidance or None,
            ),
        )
        step2 = step2_result.get("verified_translation", step1)
        step3_result = await call_for_step(
            3,
            lambda c: c.literary_revise(
                source_text, step2, step2_result, translation.source_lang, translation.target_lang,
                translation.literary_type, guidance=guidance or None,
            ),
        )
        step3 = step3_result.get("revised_translation", step2)
        step4_result = await call_for_step(
            4,
            lambda c: c.literary_finalize(
                source_text, step3, translation.source_lang, translation.target_lang,
                translation.literary_type, guidance=guidance or None,
            ),
        )
        step4 = step4_result.get("final_translation", step3)
        errata_result = await call_for_step(
            4,
            lambda c: c.literary_holistic_errata(
                source_text, step4, translation.source_lang, translation.target_lang,
                translation.literary_type, guidance=guidance or None, paragraph_count=1,
            ),
        )
        step4 = AIClient._strip_translation_wrappers(
            errata_result.get("corrected_translation", step4) or step4
        )
        result = {
            "step1_translation": step1,
            "step2_verification": step2,
            "step3_revision": step3,
            "step4_finalization": step4,
            "beauty_scores": {
                "sound": step2_result.get("beauty_sound_score", 7.0),
                "word": step2_result.get("beauty_word_score", 7.0),
                "meaning": step2_result.get("beauty_meaning_score", 7.0),
            },
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"重译失败: {e}") from e
    finally:
        reset_collab_mode(mode_token)

    paragraph.step1_translation = result["step1_translation"]
    paragraph.step2_verification = result["step2_verification"]
    paragraph.step3_revision = result["step3_revision"]
    paragraph.step4_finalization = result["step4_finalization"]
    paragraph.translated_text = result["step4_finalization"]
    paragraph.user_edited_text = None
    paragraph.is_edited = False
    beauty_scores = result.get("beauty_scores") or {}
    paragraph.beauty_sound_score = beauty_scores.get("sound", 7.0)
    paragraph.beauty_word_score = beauty_scores.get("word", 7.0)
    paragraph.beauty_meaning_score = beauty_scores.get("meaning", 7.0)
    db.commit()

    # 刷新全文聚合字段
    all_paras = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation.id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    translation.step1_translation = "\n\n".join(p.step1_translation or "" for p in all_paras)
    translation.step2_verification = "\n\n".join(p.step2_verification or "" for p in all_paras)
    translation.step3_revision = "\n\n".join(p.step3_revision or "" for p in all_paras)
    translation.step4_finalization = "\n\n".join(p.step4_finalization or "" for p in all_paras)
    translation.final_translation = "\n\n".join(
        p.user_edited_text or p.translated_text or "" for p in all_paras
    )
    scored = [p for p in all_paras if p.beauty_sound_score is not None]
    if scored:
        translation.beauty_sound_score = sum(p.beauty_sound_score or 0 for p in scored) / len(scored)
        translation.beauty_word_score = sum(p.beauty_word_score or 0 for p in scored) / len(scored)
        translation.beauty_meaning_score = sum(p.beauty_meaning_score or 0 for p in scored) / len(scored)
    db.commit()
    db.refresh(paragraph)
    return paragraph


# ============================================================
# 导出功能 - 支持多种格式
# ============================================================

def generate_txt_content(translation, paragraphs, include_source: bool = False) -> str:
    """生成纯文本格式内容"""
    lines = []
    
    if translation.title:
        lines.append(f"{translation.title}")
        lines.append("=" * 50)
        lines.append("")
    
    lines.append(f"原文语言: {translation.source_lang}")
    lines.append(f"译文语言: {translation.target_lang}")
    lines.append(f"文学类型: {translation.literary_type}")
    lines.append(f"生成时间: {_safe_datetime_str(translation.created_at)}")
    lines.append("")
    
    if translation.beauty_sound_score:
        lines.append(f"音美评分: {translation.beauty_sound_score:.1f}/10")
        lines.append(f"词美评分: {translation.beauty_word_score:.1f}/10")
        lines.append(f"意美评分: {translation.beauty_meaning_score:.1f}/10")
        lines.append("")
    
    lines.append("=" * 50)
    lines.append("")
    
    for i, para in enumerate(paragraphs, 1):
        if include_source:
            lines.append(f"【原文 {i}】")
            lines.append(para.source_text)
            lines.append("")
            lines.append(f"【译文 {i}】")
        
        text = para.user_edited_text or para.translated_text or ""
        lines.append(text)
        lines.append("")
    
    return "\n".join(lines)


def generate_md_content(translation, paragraphs, include_source: bool = False) -> str:
    """生成 Markdown 格式内容"""
    lines = []
    
    if translation.title:
        lines.append(f"# {translation.title}")
        lines.append("")
    else:
        lines.append("# 文学翻译结果")
        lines.append("")
    
    lines.append("## 翻译信息")
    lines.append("")
    lines.append(f"- **原文语言**: {translation.source_lang}")
    lines.append(f"- **译文语言**: {translation.target_lang}")
    lines.append(f"- **文学类型**: {translation.literary_type}")
    lines.append(f"- **生成时间**: {_safe_datetime_str(translation.created_at)}")
    lines.append("")
    
    if translation.beauty_sound_score:
        lines.append("## 三美评分")
        lines.append("")
        lines.append(f"- **音美**: {translation.beauty_sound_score:.1f}/10")
        lines.append(f"- **词美**: {translation.beauty_word_score:.1f}/10")
        lines.append(f"- **意美**: {translation.beauty_meaning_score:.1f}/10")
        lines.append("")
    
    lines.append("---")
    lines.append("")
    lines.append("## 翻译内容")
    lines.append("")
    
    for i, para in enumerate(paragraphs, 1):
        if include_source:
            lines.append(f"### 段落 {i}")
            lines.append("")
            lines.append("**原文：**")
            lines.append("")
            lines.append(f"> {para.source_text}")
            lines.append("")
            lines.append("**译文：**")
            lines.append("")
        
        text = para.user_edited_text or para.translated_text or ""
        lines.append(text)
        lines.append("")
        if include_source:
            lines.append("---")
            lines.append("")
    
    return "\n".join(lines)


def generate_html_content(translation, paragraphs, include_source: bool = False) -> str:
    """生成 HTML 格式内容"""
    html_parts = []
    
    html_parts.append("<!DOCTYPE html>")
    html_parts.append("<html lang=\"zh-CN\">")
    html_parts.append("<head>")
    html_parts.append("<meta charset=\"UTF-8\">")
    html_parts.append(f"<title>{translation.title or '文学翻译结果'}</title>")
    html_parts.append("<style>")
    html_parts.append("""
        body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; max-width: 800px; margin: 0 auto; padding: 40px 20px; line-height: 1.8; color: #333; }
        h1 { color: #2c3e50; border-bottom: 3px solid #3498db; padding-bottom: 10px; }
        h2 { color: #34495e; margin-top: 30px; }
        .info-box { background: #f8f9fa; padding: 20px; border-radius: 8px; margin: 20px 0; }
        .scores { display: flex; gap: 20px; margin: 10px 0; }
        .score-item { background: #e8f4f8; padding: 10px 20px; border-radius: 6px; }
        .paragraph { margin: 30px 0; padding: 20px; background: #fff; border: 1px solid #e1e8ed; border-radius: 8px; }
        .source { color: #666; font-style: italic; margin-bottom: 15px; padding-bottom: 15px; border-bottom: 1px dashed #ddd; }
        .translation { color: #2c3e50; font-size: 1.1em; }
        .para-num { color: #3498db; font-weight: bold; margin-bottom: 10px; }
        hr { border: none; border-top: 2px solid #ecf0f1; margin: 30px 0; }
    """)
    html_parts.append("</style>")
    html_parts.append("</head>")
    html_parts.append("<body>")
    
    # 标题
    html_parts.append(f"<h1>{translation.title or '文学翻译结果'}</h1>")
    
    # 信息区域
    html_parts.append("<div class='info-box'>")
    html_parts.append(f"<p><strong>原文语言：</strong>{translation.source_lang}</p>")
    html_parts.append(f"<p><strong>译文语言：</strong>{translation.target_lang}</p>")
    html_parts.append(f"<p><strong>文学类型：</strong>{translation.literary_type}</p>")
    html_parts.append(f"<p><strong>生成时间：</strong>{_safe_datetime_str(translation.created_at)}</p>")
    
    if translation.beauty_sound_score:
        html_parts.append("<div class='scores'>")
        html_parts.append(f"<div class='score-item'><strong>音美:</strong> {translation.beauty_sound_score:.1f}/10</div>")
        html_parts.append(f"<div class='score-item'><strong>词美:</strong> {translation.beauty_word_score:.1f}/10</div>")
        html_parts.append(f"<div class='score-item'><strong>意美:</strong> {translation.beauty_meaning_score:.1f}/10</div>")
        html_parts.append("</div>")
    
    html_parts.append("</div>")
    html_parts.append("<hr>")
    
    # 内容区域
    for i, para in enumerate(paragraphs, 1):
        html_parts.append("<div class='paragraph'>")
        html_parts.append(f"<div class='para-num'>段落 {i}</div>")
        
        if include_source:
            html_parts.append(f"<div class='source'>{para.source_text}</div>")
        
        text = para.user_edited_text or para.translated_text or ""
        # 处理换行
        text_html = text.replace("\n", "<br>")
        html_parts.append(f"<div class='translation'>{text_html}</div>")
        html_parts.append("</div>")
    
    html_parts.append("</body>")
    html_parts.append("</html>")
    
    return "\n".join(html_parts)


def generate_json_content(translation, paragraphs, include_source: bool = False) -> str:
    """生成 JSON 格式内容"""
    import json
    
    data = {
        "title": translation.title,
        "source_lang": translation.source_lang,
        "target_lang": translation.target_lang,
        "literary_type": translation.literary_type,
        "created_at": translation.created_at.isoformat() if translation.created_at else None,
        "beauty_scores": {
            "sound": float(translation.beauty_sound_score) if translation.beauty_sound_score else None,
            "word": float(translation.beauty_word_score) if translation.beauty_word_score else None,
            "meaning": float(translation.beauty_meaning_score) if translation.beauty_meaning_score else None
        } if translation.beauty_sound_score else None,
        "paragraphs": []
    }
    
    for i, para in enumerate(paragraphs, 1):
        para_data = {
            "index": i,
            "translation": para.user_edited_text or para.translated_text or ""
        }
        if include_source:
            para_data["source"] = para.source_text
        data["paragraphs"].append(para_data)
    
    return json.dumps(data, ensure_ascii=False, indent=2)


def generate_csv_content(translation, paragraphs, include_source: bool = False) -> str:
    """生成 CSV 格式内容"""
    import csv
    import io
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    # 写入头部信息
    writer.writerow(["文学翻译结果"])
    writer.writerow(["标题", translation.title or ""])
    writer.writerow(["原文语言", translation.source_lang])
    writer.writerow(["译文语言", translation.target_lang])
    writer.writerow(["文学类型", translation.literary_type])
    writer.writerow(["生成时间", _safe_datetime_str(translation.created_at)])
    
    if translation.beauty_sound_score:
        writer.writerow(["音美评分", f"{translation.beauty_sound_score:.1f}"])
        writer.writerow(["词美评分", f"{translation.beauty_word_score:.1f}"])
        writer.writerow(["意美评分", f"{translation.beauty_meaning_score:.1f}"])
    
    writer.writerow([])  # 空行
    
    # 写入段落内容
    if include_source:
        writer.writerow(["段落编号", "原文", "译文"])
        for i, para in enumerate(paragraphs, 1):
            text = para.user_edited_text or para.translated_text or ""
            writer.writerow([i, para.source_text, text])
    else:
        writer.writerow(["段落编号", "译文"])
        for i, para in enumerate(paragraphs, 1):
            text = para.user_edited_text or para.translated_text or ""
            writer.writerow([i, text])
    
    return output.getvalue()


@router.post("/translations/{translation_id}/export")
async def export_translation(
    translation_id: int,
    request: ExportTranslationRequest,
    db: Session = Depends(get_db)
):
    """导出翻译结果 - 支持多种格式: txt, md, html, json, csv"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    
    format_generators = {
        "txt": generate_txt_content,
        "md": generate_md_content,
        "html": generate_html_content,
        "json": generate_json_content,
        "csv": generate_csv_content,
    }
    
    format_mime_types = {
        "txt": "text/plain; charset=utf-8",
        "md": "text/markdown; charset=utf-8",
        "html": "text/html; charset=utf-8",
        "json": "application/json; charset=utf-8",
        "csv": "text/csv; charset=utf-8",
    }
    
    format_ext = request.format.lower()
    
    if format_ext not in format_generators:
        raise HTTPException(
            status_code=400, 
            detail=f"不支持的格式: {request.format}. 支持的格式: {', '.join(format_generators.keys())}"
        )
    
    # 生成内容
    content = format_generators[format_ext](
        translation,
        paragraphs,
        request.include_source
    )

    raw_name = (translation.title or "").strip() or f"literary_translation_{translation_id}"
    # Windows / 通用文件名非法字符
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", raw_name).strip(" .")
    safe_name = re.sub(r"\s+", " ", safe_name) or f"literary_translation_{translation_id}"
    if len(safe_name) > 120:
        safe_name = safe_name[:120].rstrip(" .")

    return {
        "filename": f"{safe_name}.{format_ext}",
        "content": content,
        "format": request.format,
        "mime_type": format_mime_types[format_ext]
    }


# ============================================================
# 专业词库管理
# ============================================================

@router.get("/terms", response_model=List[ProfessionalTermResponse])
async def list_professional_terms(
    literary_type: Optional[str] = None,
    category: Optional[str] = None,
    source_lang: Optional[str] = None,
    target_lang: Optional[str] = None,
    keyword: Optional[str] = None,
    is_verified: Optional[bool] = None,
    translation_id: Optional[int] = None,
    scope: Optional[str] = Query(None, description="document=小词库, global=大词库"),
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    """获取专业词汇。scope=document 需 translation_id；scope=global 为系统大词库。"""
    query = db.query(ProfessionalTerm)

    if scope == "document":
        if not translation_id:
            raise HTTPException(status_code=400, detail="小词库需要 translation_id")
        query = query.filter(ProfessionalTerm.translation_id == translation_id)
    elif scope == "global":
        query = query.filter(ProfessionalTerm.translation_id.is_(None))
    elif translation_id is not None:
        query = query.filter(ProfessionalTerm.translation_id == translation_id)

    if literary_type:
        query = query.filter(ProfessionalTerm.literary_type == literary_type)
    if category:
        query = query.filter(ProfessionalTerm.category == category)
    if source_lang:
        query = query.filter(ProfessionalTerm.source_lang == source_lang)
    if target_lang:
        query = query.filter(ProfessionalTerm.target_lang == target_lang)
    if is_verified is not None:
        query = query.filter(ProfessionalTerm.is_verified == is_verified)
    if keyword:
        query = query.filter(
            (ProfessionalTerm.source_term.contains(keyword)) |
            (ProfessionalTerm.target_term.contains(keyword))
        )

    query = query.order_by(ProfessionalTerm.usage_count.desc(), ProfessionalTerm.created_at.desc())
    return query.offset(skip).limit(limit).all()


def _upsert_global_term(
    db: Session,
    *,
    source_term: str,
    target_term: str,
    literary_type: str,
    source_lang: str,
    target_lang: str,
    category: Optional[str] = None,
    description: Optional[str] = None,
    usage: int = 1,
    cache: Optional[dict] = None,
) -> str:
    """把一条词汇归入系统大词典（translation_id 为空）。同分类同语言对只保留一条。"""
    source_term = (source_term or "").strip()
    target_term = (target_term or "").strip()
    literary_type = literary_type or "general"
    if not source_term or not target_term:
        return "skipped"
    key = (source_term, literary_type, source_lang, target_lang)
    existing = cache.get(key) if cache is not None else None
    if existing is None:
        existing = db.query(ProfessionalTerm).filter(
            ProfessionalTerm.source_term == source_term,
            ProfessionalTerm.literary_type == literary_type,
            ProfessionalTerm.source_lang == source_lang,
            ProfessionalTerm.target_lang == target_lang,
            ProfessionalTerm.translation_id.is_(None),
        ).first()
    if existing:
        if category and not existing.category:
            existing.category = category
        if description and not existing.description:
            existing.description = description
        existing.usage_count = (existing.usage_count or 0) + max(usage or 1, 1)
        if cache is not None:
            cache[key] = existing
        return "updated"
    created = ProfessionalTerm(
        source_term=source_term,
        target_term=target_term,
        literary_type=literary_type,
        category=category,
        source_lang=source_lang,
        target_lang=target_lang,
        description=description,
        translation_id=None,
        usage_count=max(usage or 1, 1),
        is_verified=True,
    )
    db.add(created)
    if cache is not None:
        cache[key] = created
    return "created"


def sync_document_terms_into_global(db: Session) -> dict:
    """把各项目小词典词汇按文学类型归并进系统大词典，并写入数据库。小词典条目保留。"""
    rows = db.query(ProfessionalTerm, LiteraryTranslation.literary_type).outerjoin(
        LiteraryTranslation, ProfessionalTerm.translation_id == LiteraryTranslation.id
    ).filter(ProfessionalTerm.translation_id.isnot(None)).all()
    created = updated = skipped = 0
    cache: dict = {}
    for term, parent_type in rows:
        result = _upsert_global_term(
            db,
            source_term=term.source_term,
            target_term=term.target_term,
            literary_type=term.literary_type or parent_type or "general",
            source_lang=term.source_lang,
            target_lang=term.target_lang,
            category=term.category,
            description=term.description,
            usage=term.usage_count or 1,
            cache=cache,
        )
        if result == "created":
            created += 1
        elif result == "updated":
            updated += 1
        else:
            skipped += 1
    db.commit()
    return {
        "scanned": len(rows),
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "global_total": db.query(ProfessionalTerm).filter(ProfessionalTerm.translation_id.is_(None)).count(),
    }


@router.post("/terms/sync-global")
async def sync_terms_into_global_dictionary(db: Session = Depends(get_db)):
    """将全部项目小词典归类写入系统大词典。"""
    return sync_document_terms_into_global(db)


@router.post("/terms", response_model=ProfessionalTermResponse)
async def create_professional_term(
    request: ProfessionalTermCreate,
    db: Session = Depends(get_db)
):
    """创建专业词汇。带 translation_id 写入小词库，否则写入大词库。"""
    literary_type = request.literary_type.value if hasattr(request.literary_type, 'value') else request.literary_type
    query = db.query(ProfessionalTerm).filter(
        ProfessionalTerm.source_term == request.source_term,
        ProfessionalTerm.literary_type == literary_type,
        ProfessionalTerm.source_lang == request.source_lang,
        ProfessionalTerm.target_lang == request.target_lang,
    )
    if request.translation_id is not None:
        query = query.filter(ProfessionalTerm.translation_id == request.translation_id)
    else:
        query = query.filter(ProfessionalTerm.translation_id.is_(None))
    existing = query.first()

    if existing:
        existing.target_term = request.target_term
        existing.category = request.category
        existing.description = request.description
        existing.usage_count += 1
        db.commit()
        db.refresh(existing)
        return existing

    if request.translation_id is not None:
        exists = db.query(LiteraryTranslation.id).filter(
            LiteraryTranslation.id == request.translation_id
        ).first()
        if not exists:
            raise HTTPException(status_code=404, detail="Translation not found")

    term = ProfessionalTerm(
        source_term=request.source_term,
        target_term=request.target_term,
        literary_type=literary_type,
        category=request.category,
        source_lang=request.source_lang,
        target_lang=request.target_lang,
        description=request.description,
        translation_id=request.translation_id,
        usage_count=1,
        is_verified=True
    )
    db.add(term)
    db.commit()
    db.refresh(term)
    if request.translation_id is not None:
        _upsert_global_term(
            db,
            source_term=term.source_term,
            target_term=term.target_term,
            literary_type=literary_type,
            source_lang=term.source_lang,
            target_lang=term.target_lang,
            category=term.category,
            description=term.description,
            usage=1,
        )
        db.commit()
    return term


@router.put("/terms/{term_id}", response_model=ProfessionalTermResponse)
async def update_professional_term(
    term_id: int,
    request: ProfessionalTermUpdate,
    db: Session = Depends(get_db)
):
    """更新专业词汇"""
    term = db.query(ProfessionalTerm).filter(ProfessionalTerm.id == term_id).first()
    if not term:
        raise HTTPException(status_code=404, detail="Term not found")

    if request.target_term is not None:
        term.target_term = request.target_term
    if request.category is not None:
        term.category = request.category
    if request.description is not None:
        term.description = request.description
    if request.is_verified is not None:
        term.is_verified = request.is_verified
    if "translation_id" in request.model_fields_set:
        term.translation_id = request.translation_id

    db.commit()
    db.refresh(term)
    return term


@router.post("/terms/{term_id}/promote")
async def promote_term_to_global(
    term_id: int,
    db: Session = Depends(get_db)
):
    """将小词库词汇提升到系统大词库"""
    term = db.query(ProfessionalTerm).filter(ProfessionalTerm.id == term_id).first()
    if not term:
        raise HTTPException(status_code=404, detail="Term not found")
    if term.translation_id is None:
        return {"message": "已在大词库", "term_id": term.id}

    conflict = db.query(ProfessionalTerm).filter(
        ProfessionalTerm.source_term == term.source_term,
        ProfessionalTerm.literary_type == term.literary_type,
        ProfessionalTerm.source_lang == term.source_lang,
        ProfessionalTerm.target_lang == term.target_lang,
        ProfessionalTerm.translation_id.is_(None),
        ProfessionalTerm.id != term.id,
    ).first()
    if conflict:
        conflict.target_term = term.target_term
        conflict.category = term.category or conflict.category
        conflict.description = term.description or conflict.description
        conflict.usage_count += term.usage_count or 1
        db.delete(term)
        db.commit()
        return {"message": "已合并到大词库", "term_id": conflict.id}

    term.translation_id = None
    db.commit()
    return {"message": "已提升到大词库", "term_id": term.id}


@router.delete("/terms/{term_id}")
async def delete_professional_term(
    term_id: int,
    db: Session = Depends(get_db)
):
    """删除专业词汇"""
    term = db.query(ProfessionalTerm).filter(ProfessionalTerm.id == term_id).first()
    if not term:
        raise HTTPException(status_code=404, detail="Term not found")

    db.delete(term)
    db.commit()
    return {"message": "Term deleted successfully"}


@router.get("/terms/categories")
async def get_term_categories(
    literary_type: Optional[str] = None,
    translation_id: Optional[int] = None,
    scope: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """获取词汇分类列表"""
    query = db.query(ProfessionalTerm.category).distinct()
    if scope == "document" and translation_id:
        query = query.filter(ProfessionalTerm.translation_id == translation_id)
    elif scope == "global":
        query = query.filter(ProfessionalTerm.translation_id.is_(None))
    if literary_type:
        query = query.filter(ProfessionalTerm.literary_type == literary_type)

    categories = [c[0] for c in query.all() if c[0]]
    return categories


# ============================================================
# 翻译词汇总结
# ============================================================

@router.get("/translations/{translation_id}/terms", response_model=TranslationTermSummaryResponse)
async def get_translation_term_summary(
    translation_id: int,
    db: Session = Depends(get_db)
):
    """获取翻译任务的专业词汇总结"""
    summary = db.query(TranslationTermSummary).filter(
        TranslationTermSummary.translation_id == translation_id
    ).first()

    if not summary:
        raise HTTPException(status_code=404, detail="Term summary not found")

    return summary


@router.post("/translations/{translation_id}/terms/extract")
async def extract_terms_from_translation(
    translation_id: int,
    db: Session = Depends(get_db)
):
    """从翻译任务中提取专业词汇"""
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()

    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")

    # 获取该类型的已有词汇
    existing_terms = db.query(ProfessionalTerm).filter(
        ProfessionalTerm.literary_type == translation.literary_type,
        ProfessionalTerm.source_lang == translation.source_lang,
        ProfessionalTerm.target_lang == translation.target_lang
    ).all()

    existing_terms_list = [
        {"source_term": t.source_term, "target_term": t.target_term}
        for t in existing_terms
    ]

    # 使用 AI 提取词汇
    client = get_ai_client()
    result = await client.extract_professional_terms(
        source_text=translation.source_text,
        translated_text=translation.final_translation or translation.step4_finalization or "",
        source_lang=translation.source_lang,
        target_lang=translation.target_lang,
        literary_type=translation.literary_type,
        existing_terms=existing_terms_list
    )

    # 处理提取的词汇
    new_count = 0
    updated_count = 0
    terms_for_summary = []

    for term_data in result.get("terms", []):
        source_term = term_data.get("source_term", "")
        target_term = term_data.get("target_term", "")
        category = term_data.get("category", "")
        description = term_data.get("description", "")

        if not source_term or not target_term:
            continue

        terms_for_summary.append({
            "source_term": source_term,
            "target_term": target_term,
            "category": category,
            "description": description
        })

        # 检查是否已存在
        existing = db.query(ProfessionalTerm).filter(
            ProfessionalTerm.source_term == source_term,
            ProfessionalTerm.literary_type == translation.literary_type,
            ProfessionalTerm.source_lang == translation.source_lang,
            ProfessionalTerm.target_lang == translation.target_lang
        ).first()

        if existing:
            # 更新现有词汇
            existing.target_term = target_term
            existing.category = category
            existing.description = description
            existing.usage_count += 1
            updated_count += 1
        else:
            # 创建新词汇
            new_term = ProfessionalTerm(
                source_term=source_term,
                target_term=target_term,
                literary_type=translation.literary_type,
                category=category,
                source_lang=translation.source_lang,
                target_lang=translation.target_lang,
                description=description,
                usage_count=1,
                translation_id=translation_id,
                is_verified=True
            )
            db.add(new_term)
            new_count += 1

    db.commit()
    sync_document_terms_into_global(db)

    # 创建总结记录
    summary = TranslationTermSummary(
        translation_id=translation_id,
        terms=terms_for_summary,
        total_terms=len(terms_for_summary),
        new_terms=new_count,
        updated_terms=updated_count,
        summary_text=result.get("summary", "")
    )
    db.add(summary)
    db.commit()
    db.refresh(summary)

    return {
        "message": f"提取完成，新增 {new_count} 个词汇，更新 {updated_count} 个词汇",
        "summary": summary
    }


# ============================================================
# 辅助函数：定稿时自动提取词汇
# ============================================================

async def auto_extract_terms_after_finalize(
    translation_id: int,
    db: Session
):
    """定稿后自动提取专业词汇"""
    try:
        translation = db.query(LiteraryTranslation).filter(
            LiteraryTranslation.id == translation_id
        ).first()

        if not translation or not (translation.final_translation or translation.step4_finalization):
            return

        # 获取该类型的已有词汇
        existing_terms = db.query(ProfessionalTerm).filter(
            ProfessionalTerm.literary_type == translation.literary_type,
            ProfessionalTerm.source_lang == translation.source_lang,
            ProfessionalTerm.target_lang == translation.target_lang
        ).all()

        existing_terms_list = [
            {"source_term": t.source_term, "target_term": t.target_term}
            for t in existing_terms
        ]

        # 使用 AI 提取词汇
        client = get_ai_client()
        result = await client.extract_professional_terms(
            source_text=translation.source_text,
            translated_text=translation.final_translation or translation.step4_finalization,
            source_lang=translation.source_lang,
            target_lang=translation.target_lang,
            literary_type=translation.literary_type,
            existing_terms=existing_terms_list
        )

        # 处理提取的词汇
        new_count = 0
        updated_count = 0
        terms_for_summary = []

        for term_data in result.get("terms", []):
            source_term = term_data.get("source_term", "")
            target_term = term_data.get("target_term", "")
            category = term_data.get("category", "")
            description = term_data.get("description", "")

            if not source_term or not target_term:
                continue

            terms_for_summary.append({
                "source_term": source_term,
                "target_term": target_term,
                "category": category,
                "description": description
            })

            # 检查是否已存在
            existing = db.query(ProfessionalTerm).filter(
                ProfessionalTerm.source_term == source_term,
                ProfessionalTerm.literary_type == translation.literary_type,
                ProfessionalTerm.source_lang == translation.source_lang,
                ProfessionalTerm.target_lang == translation.target_lang
            ).first()

            if existing:
                existing.target_term = target_term
                existing.category = category
                existing.description = description
                existing.usage_count += 1
                updated_count += 1
            else:
                new_term = ProfessionalTerm(
                    source_term=source_term,
                    target_term=target_term,
                    literary_type=translation.literary_type,
                    category=category,
                    source_lang=translation.source_lang,
                    target_lang=translation.target_lang,
                    description=description,
                    usage_count=1,
                    translation_id=translation_id,
                    is_verified=True
                )
                db.add(new_term)
                new_count += 1

        db.commit()
        sync_document_terms_into_global(db)

        # 创建总结记录
        summary = TranslationTermSummary(
            translation_id=translation_id,
            terms=terms_for_summary,
            total_terms=len(terms_for_summary),
            new_terms=new_count,
            updated_terms=updated_count,
            summary_text=result.get("summary", "")
        )
        db.add(summary)
        db.commit()

    except Exception as e:
        print(f"[Auto Extract Terms] Error: {e}")
        # 不抛出异常，避免影响定稿流程


# ============================================================
# 长文件一键翻译
# ============================================================

# 支持的上传格式：txt、docx、pdf、mobi 及常见文本格式
SUPPORTED_UPLOAD_EXTENSIONS = {
    'txt', 'md', 'markdown', 'text',
    'docx', 'doc',   # Word
    'pdf',
    'mobi', 'azw', 'azw3',  # 电子书（含 KF8）
    'html', 'htm', 'xhtml', 'xml', 'json', 'csv',
    'log', 'rst', 'tex', 'srt', 'sub', 'vtt', 'yaml', 'yml', 'ini', 'cfg', 'properties',
}


def _decode_text_content(content: bytes) -> str:
    """多种编码尝试解码"""
    for encoding in ('utf-8', 'utf-8-sig', 'gbk', 'gb2312', 'latin-1'):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode('utf-8', errors='ignore')


def _extract_text_docx(content: bytes) -> str:
    """从 Word docx 提取正文"""
    from docx import Document
    doc = Document(io.BytesIO(content))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_text_pdf(content: bytes) -> str:
    """从 PDF 提取正文"""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            parts.append(t)
    return "\n\n".join(parts)


def _drop_html_block(raw: str, tag: str) -> str:
    """线性去掉 script/style，避免大文件上正则回溯卡住。"""
    low = raw.lower()
    open_mark = f"<{tag}"
    close_mark = f"</{tag}>"
    out: List[str] = []
    i = 0
    n = len(raw)
    while i < n:
        start = low.find(open_mark, i)
        if start < 0:
            out.append(raw[i:])
            break
        out.append(raw[i:start])
        gt = low.find(">", start)
        if gt < 0:
            break
        end = low.find(close_mark, gt)
        if end < 0:
            i = gt + 1
            continue
        i = end + len(close_mark)
    return "".join(out)


def _html_to_plain_text(raw: str) -> str:
    """HTML/XHTML 转纯文本，保留段落换行。"""
    import html as html_lib
    text = _drop_html_block(raw, "script")
    text = _drop_html_block(text, "style")
    text = re.sub(r"(?i)<\s*br\s*/?\s*>", "\n", text)
    text = re.sub(
        r"(?i)</\s*(p|div|h[1-6]|li|tr|section|article|blockquote)\s*>",
        "\n\n",
        text,
    )
    text = re.sub(r"<[^>]+>", " ", text)
    text = html_lib.unescape(text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def _readable_score(text: str) -> int:
    if not text:
        return 0
    return len(re.findall(r"[A-Za-z0-9\u4e00-\u9fff]", text))


def _extract_text_from_epub(epub_path: str) -> str:
    """从 EPUB（ZIP）中按阅读顺序提取 HTML/XHTML 正文。"""
    import zipfile
    from xml.etree import ElementTree as ET

    def _local(tag: str) -> str:
        return tag.split("}")[-1] if "}" in tag else tag

    with zipfile.ZipFile(epub_path, "r") as zf:
        names = zf.namelist()
        # 定位 OPF
        opf_path = None
        if "META-INF/container.xml" in names:
            try:
                root = ET.fromstring(zf.read("META-INF/container.xml"))
                for node in root.iter():
                    if _local(node.tag) == "rootfile":
                        opf_path = node.attrib.get("full-path")
                        break
            except ET.ParseError:
                opf_path = None
        if not opf_path:
            for name in names:
                if name.lower().endswith(".opf"):
                    opf_path = name
                    break

        spine_hrefs: List[str] = []
        if opf_path and opf_path in names:
            try:
                opf = ET.fromstring(zf.read(opf_path))
                manifest = {}
                for node in opf.iter():
                    if _local(node.tag) == "item":
                        item_id = node.attrib.get("id")
                        href = node.attrib.get("href")
                        if item_id and href:
                            manifest[item_id] = href
                for node in opf.iter():
                    if _local(node.tag) == "itemref":
                        idref = node.attrib.get("idref")
                        href = manifest.get(idref or "")
                        if href:
                            spine_hrefs.append(href)
            except ET.ParseError:
                spine_hrefs = []

        opf_dir = opf_path.rsplit("/", 1)[0] if opf_path and "/" in opf_path else ""

        def _resolve(href: str) -> Optional[str]:
            from urllib.parse import unquote
            href = unquote((href or "").split("#", 1)[0].split("?", 1)[0]).replace("\\", "/")
            candidates = []
            if opf_dir:
                candidates.append(f"{opf_dir}/{href}".replace("\\", "/"))
            candidates.append(href)
            # 有些路径带 ../
            for c in candidates:
                parts = []
                for part in c.split("/"):
                    if part in ("", "."):
                        continue
                    if part == "..":
                        if parts:
                            parts.pop()
                        continue
                    parts.append(part)
                norm = "/".join(parts)
                if norm in names:
                    return norm
            # 退化为 basename 匹配
            base = href.split("/")[-1]
            for name in names:
                if name.endswith("/" + base) or name == base:
                    return name
            return None

        html_parts: List[str] = []
        if spine_hrefs:
            for href in spine_hrefs:
                path = _resolve(href)
                if not path:
                    continue
                raw = zf.read(path).decode("utf-8", errors="ignore")
                part = _html_to_plain_text(raw)
                if part:
                    html_parts.append(part)
        if not html_parts:
            # 无 spine 时扫全部 html/xhtml
            for name in names:
                lower = name.lower()
                if lower.endswith((".html", ".xhtml", ".htm")) and "meta-inf" not in lower:
                    raw = zf.read(name).decode("utf-8", errors="ignore")
                    part = _html_to_plain_text(raw)
                    if part:
                        html_parts.append(part)
        return "\n\n".join(html_parts).strip()


def _long_win_path(path: str) -> str:
    ap = os.path.abspath(path)
    if ap.startswith("\\\\?\\"):
        return ap
    return "\\\\?\\" + ap


def _open_tolerant(path, mode="r", **kwargs):
    """Windows 上超长路径或非法参数时改用 \\\\?\\ 前缀再打开。"""
    try:
        return open(path, mode, **kwargs)
    except OSError as e:
        if os.name == "nt" and getattr(e, "errno", None) == 22 and isinstance(path, str):
            return open(_long_win_path(path), mode, **kwargs)
        raise


def _read_html_file(path: str) -> str:
    with _open_tolerant(path, "r", encoding="utf-8", errors="ignore") as f:
        return _html_to_plain_text(f.read())


def _extract_text_mobi(content: bytes) -> str:
    """从 MOBI/AZW/AZW3 提取正文。

    以 mobi.extract 返回的文件为主（原先能解析的书走这条）。
    双格式书还会额外尝试 EPUB 目录和 mobi7/book.html，取可读文字更多的一份。
    某一路失败不会把整本判失败。
    """
    import builtins
    import contextlib
    import io
    import mobi

    @contextlib.contextmanager
    def _unpack_compat():
        """解包时屏蔽 print（Windows 控制台会抛 Errno 22），并重试超长路径。"""
        orig_open = builtins.open
        orig_mkdir = os.mkdir

        def safe_open(file, *args, **kwargs):
            try:
                return orig_open(file, *args, **kwargs)
            except OSError as e:
                if (
                    os.name == "nt"
                    and getattr(e, "errno", None) == 22
                    and isinstance(file, (str, os.PathLike))
                ):
                    return orig_open(_long_win_path(os.fspath(file)), *args, **kwargs)
                raise

        def safe_mkdir(dir_path, *args, **kwargs):
            try:
                return orig_mkdir(dir_path, *args, **kwargs)
            except OSError as e:
                if (
                    os.name == "nt"
                    and getattr(e, "errno", None) == 22
                    and isinstance(dir_path, (str, os.PathLike))
                ):
                    return orig_mkdir(_long_win_path(os.fspath(dir_path)), *args, **kwargs)
                raise

        builtins.open = safe_open
        os.mkdir = safe_mkdir
        sink = io.StringIO()
        try:
            with contextlib.redirect_stdout(sink), contextlib.redirect_stderr(sink):
                yield
        finally:
            builtins.open = orig_open
            os.mkdir = orig_mkdir

    fd, path = tempfile.mkstemp(suffix=".mobi")
    try:
        os.write(fd, content)
        os.close(fd)
        with _unpack_compat():
            tempdir, filepath = mobi.extract(path)
        try:
            ext = filepath.split(".")[-1].lower() if "." in filepath else ""
            candidates: List[str] = []

            if ext == "pdf":
                try:
                    with open(filepath, "rb") as f:
                        candidates.append(_extract_text_pdf(f.read()) or "")
                except Exception:
                    pass
            else:
                try:
                    candidates.append(_read_html_file(filepath))
                except Exception:
                    pass

            if ext == "epub":
                try:
                    candidates.append(_extract_text_from_epub(filepath) or "")
                except Exception:
                    pass

            html_fallback = os.path.join(tempdir, "mobi7", "book.html")
            if os.path.isfile(html_fallback) and os.path.abspath(html_fallback) != os.path.abspath(filepath):
                try:
                    candidates.append(_read_html_file(html_fallback))
                except Exception:
                    pass

            best = max(candidates, key=_readable_score, default="")
            if best and best.strip():
                return best.strip()
            raise ValueError("未提取到正文（可能是 DRM 加密或损坏的电子书）")
        finally:
            shutil.rmtree(tempdir, ignore_errors=True)
    except ValueError:
        raise
    except Exception as e:
        msg = str(e) or type(e).__name__
        lower = msg.lower()
        if "drm" in lower or "encryption" in lower or "encrypted" in lower:
            raise ValueError("该电子书带有 DRM 加密，无法解析") from e
        raise ValueError(f"MOBI/AZW 解析失败: {msg}") from e
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def parse_text_file(content: bytes, file_type: str) -> str:
    """根据扩展名解析文件内容为纯文本。支持 txt、docx、pdf、mobi 等"""
    ext = (file_type or "").lower()
    try:
        if ext in ("docx", "doc"):
            return _extract_text_docx(content)
        if ext == "pdf":
            return _extract_text_pdf(content)
        if ext in ("mobi", "azw", "azw3"):
            return _extract_text_mobi(content)
        # 文本类：多编码解码
        text = _decode_text_content(content)
        return text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"无法解析文件 ({ext}): {str(e)}")


def split_long_text(text: str, max_chunk_size: int = MAX_PARAGRAPH_SIZE) -> List[str]:
    """
    将长文本分段，每段不超过 max_chunk_size 字符
    尽量在段落边界处分割；单段超长时按句子边界拆
    """
    return split_text_into_paragraphs(text, max_chunk_size)


MAX_UPLOAD_SIZE = 20 * 1024 * 1024


@router.post("/parse-file")
async def parse_uploaded_file(file: UploadFile = File(...)):
    """
    解析上传文件并返回提取的正文（不创建任务）。
    支持 txt、docx、pdf、mobi 等，用于新建任务时「上传文件」填充原文。
    """
    ext = file.filename.split('.')[-1].lower() if '.' in file.filename else ''
    if ext not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"不支持的文件格式: {ext}")
    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise HTTPException(
            status_code=413,
            detail=f"文件过大（{len(content) / 1024 / 1024:.1f}MB），最大支持 {MAX_UPLOAD_SIZE // 1024 // 1024}MB"
        )
    try:
        text = parse_text_file(content, ext)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not (text or "").strip():
        raise HTTPException(status_code=400, detail="未能从文件中提取到正文，请确认文件未加密且格式完整")
    para_count = len(split_text_into_paragraphs(text))
    return {"text": text, "filename": file.filename, "char_count": len(text), "paragraph_count": para_count}


@router.post("/translations/upload")
async def upload_and_translate_file(
    file: UploadFile = File(...),
    title: Optional[str] = Form(None),
    source_lang: str = Form("auto"),
    target_lang: str = Form("en"),
    literary_type: str = Form("general"),
    reference_document_ids: Optional[str] = Form(None),
    user_requirements: Optional[str] = Form(None),
    style_agent_id: Optional[int] = Form(None),
    auto_run: bool = Form(True, description="是否自动执行四步流程"),
    collab_mode: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    """
    上传文件并创建翻译任务。支持：txt, docx, pdf, mobi 及 md/html/xml/json/csv 等文本格式；可选自动执行四步流程。
    """
    file_extension = file.filename.split('.')[-1].lower() if '.' in file.filename else ''
    if file_extension not in SUPPORTED_UPLOAD_EXTENSIONS:
        supported = ', '.join(sorted(SUPPORTED_UPLOAD_EXTENSIONS))
        raise HTTPException(
            status_code=400,
            detail=f"不支持的文件格式: {file_extension}. 支持: {supported}"
        )
    content = await file.read()
    try:
        text_content = parse_text_file(content, file_extension)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if not text_content.strip():
        raise HTTPException(status_code=400, detail="文件内容为空")
    ref_ids = json.loads(reference_document_ids) if reference_document_ids else []
    _validate_style_agent_id(db, style_agent_id)
    translation = LiteraryTranslation(
        title=title or file.filename,
        source_text=text_content,
        source_lang=source_lang,
        target_lang=target_lang,
        literary_type=literary_type,
        reference_document_ids=ref_ids,
        user_requirements=user_requirements.strip() if user_requirements else None,
        style_agent_id=style_agent_id,
        collab_mode=_resolve_task_collab_mode(collab_mode),
        status=LiteraryTranslationStatus.PENDING,
        current_step=1
    )
    db.add(translation)
    db.commit()
    db.refresh(translation)
    chunks = split_long_text(text_content)
    for idx, chunk in enumerate(chunks):
        paragraph = LiteraryParagraph(
            translation_id=translation.id,
            paragraph_index=idx,
            source_text=chunk
        )
        db.add(paragraph)
    db.commit()
    if auto_run:
        mode_token = push_collab_mode(getattr(translation, "collab_mode", None))
        try:
            await _execute_step1(translation.id, db)
            await _execute_step2(translation.id, db)
            await _execute_step3(translation.id, db)
            await _execute_step4(translation.id, db)
            db.refresh(translation)
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"自动执行四步流程失败: {str(e)}")
        finally:
            reset_collab_mode(mode_token)
    return {
        "message": "文件上传成功" + ("，四步流程已自动执行完成" if auto_run else "，翻译任务已创建"),
        "translation_id": translation.id,
        "filename": file.filename,
        "total_chunks": len(chunks),
        "total_chars": len(text_content),
        "current_step": translation.current_step,
        "status": translation.status
    }


@router.post("/translations/upload-batch")
async def upload_and_translate_batch(
    files: List[UploadFile] = File(..., description="多个文件"),
    source_lang: str = Form("auto"),
    target_lang: str = Form("en"),
    literary_type: str = Form("general"),
    reference_document_ids: Optional[str] = Form(None),
    user_requirements: Optional[str] = Form(None),
    style_agent_id: Optional[int] = Form(None),
    auto_run: bool = Form(True, description="是否对每个任务自动执行四步流程"),
    collab_mode: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    """
    批量上传文件并创建多个翻译任务。每个文件对应一个任务；可选对每个任务自动执行四步流程。
    """
    if not files:
        raise HTTPException(status_code=400, detail="请至少上传一个文件")
    supported = ", ".join(sorted(SUPPORTED_UPLOAD_EXTENSIONS))
    results = []
    ref_ids = json.loads(reference_document_ids) if reference_document_ids else []
    req_text = user_requirements.strip() if user_requirements else None
    task_mode = _resolve_task_collab_mode(collab_mode)
    _validate_style_agent_id(db, style_agent_id)

    for file in files:
        file_extension = file.filename.split(".")[-1].lower() if "." in file.filename else ""
        if file_extension not in SUPPORTED_UPLOAD_EXTENSIONS:
            results.append({
                "filename": file.filename,
                "translation_id": None,
                "error": f"不支持的文件格式: {file_extension}. 支持: {supported}",
                "total_chunks": 0,
                "total_chars": 0,
                "current_step": 1,
                "status": "failed",
            })
            continue
        try:
            content = await file.read()
            if len(content) > MAX_UPLOAD_SIZE:
                results.append({
                    "filename": file.filename,
                    "translation_id": None,
                    "error": f"文件过大（最大 {MAX_UPLOAD_SIZE // 1024 // 1024}MB）",
                    "total_chunks": 0,
                    "total_chars": 0,
                    "current_step": 1,
                    "status": "failed",
                })
                continue
            text_content = parse_text_file(content, file_extension)
        except ValueError as e:
            results.append({
                "filename": file.filename,
                "translation_id": None,
                "error": str(e),
                "total_chunks": 0,
                "total_chars": 0,
                "current_step": 1,
                "status": "failed",
            })
            continue
        if not text_content.strip():
            results.append({
                "filename": file.filename,
                "translation_id": None,
                "error": "文件内容为空",
                "total_chunks": 0,
                "total_chars": 0,
                "current_step": 1,
                "status": "failed",
            })
            continue

        translation = LiteraryTranslation(
            title=file.filename,
            source_text=text_content,
            source_lang=source_lang,
            target_lang=target_lang,
            literary_type=literary_type,
            reference_document_ids=ref_ids,
            user_requirements=req_text,
            style_agent_id=style_agent_id,
            collab_mode=task_mode,
            status=LiteraryTranslationStatus.PENDING,
            current_step=1,
        )
        db.add(translation)
        db.commit()
        db.refresh(translation)
        chunks = split_long_text(text_content)
        for idx, chunk in enumerate(chunks):
            paragraph = LiteraryParagraph(
                translation_id=translation.id,
                paragraph_index=idx,
                source_text=chunk,
            )
            db.add(paragraph)
        db.commit()

        if auto_run:
            mode_token = push_collab_mode(getattr(translation, "collab_mode", None))
            try:
                await _execute_step1(translation.id, db)
                await _execute_step2(translation.id, db)
                await _execute_step3(translation.id, db)
                await _execute_step4(translation.id, db)
                db.refresh(translation)
            except HTTPException:
                raise
            except Exception as e:
                results.append({
                    "filename": file.filename,
                    "translation_id": translation.id,
                    "error": f"自动执行四步流程失败: {str(e)}",
                    "total_chunks": len(chunks),
                    "total_chars": len(text_content),
                    "current_step": translation.current_step,
                    "status": translation.status,
                })
                continue
            finally:
                reset_collab_mode(mode_token)

        results.append({
            "filename": file.filename,
            "translation_id": translation.id,
            "error": None,
            "total_chunks": len(chunks),
            "total_chars": len(text_content),
            "current_step": translation.current_step,
            "status": translation.status,
        })

    return {
        "message": f"批量上传完成，共 {len(files)} 个文件，成功 {sum(1 for r in results if r.get('translation_id'))} 个",
        "results": results,
    }


@router.post("/translations/{translation_id}/translate-all")
async def translate_all_chunks(
    translation_id: int,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db)
):
    """
    一键翻译所有段落（长文本批量翻译）
    """
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).order_by(LiteraryParagraph.paragraph_index).all()
    
    if not paragraphs:
        raise HTTPException(status_code=400, detail="没有需要翻译的段落")
    
    # 更新状态为翻译中
    translation.status = LiteraryTranslationStatus.TRANSLATING
    db.commit()

    reference_content = get_reference_and_requirements(translation, db)
    guidance = build_prompt_guidance(translation, db, include_story=True)
    ref_safe = reference_content if reference_content else None

    # 批量翻译所有段落
    total = len(paragraphs)
    success_count = 0
    mode_token = push_collab_mode(getattr(translation, "collab_mode", None))

    try:
        for idx, para in enumerate(paragraphs):
            try:
                text = para.source_text or ""
                step1 = await call_for_step(
                    1,
                    lambda c, src=text: c.literary_translate(
                        src, translation.source_lang, translation.target_lang,
                        translation.literary_type, ref_safe, guidance=guidance or None,
                    ),
                )
                step2_result = await call_for_step(
                    2,
                    lambda c, src=text, draft=step1: c.literary_verify(
                        src, draft, translation.source_lang, translation.target_lang,
                        translation.literary_type, guidance=guidance or None,
                    ),
                )
                step2 = step2_result.get("verified_translation", step1)
                step3_result = await call_for_step(
                    3,
                    lambda c, src=text, verified=step2, analysis=step2_result: c.literary_revise(
                        src, verified, analysis, translation.source_lang, translation.target_lang,
                        translation.literary_type, guidance=guidance or None,
                    ),
                )
                step3 = step3_result.get("revised_translation", step2)
                step4_result = await call_for_step(
                    4,
                    lambda c, src=text, revised=step3: c.literary_finalize(
                        src, revised, translation.source_lang, translation.target_lang,
                        translation.literary_type, guidance=guidance or None,
                    ),
                )
                step4 = step4_result.get("final_translation", step3)
                result = {
                    "step1_translation": step1,
                    "step2_verification": step2,
                    "step3_revision": step3,
                    "step4_finalization": step4,
                }

                # 保存四步翻译结果
                para.step1_translation = result["step1_translation"]
                para.step2_verification = result["step2_verification"]
                para.step3_revision = result["step3_revision"]
                para.step4_finalization = result["step4_finalization"]
                para.translated_text = result["step4_finalization"]

                # 保存三美评分
                beauty_scores = result.get("beauty_scores", {})
                para.beauty_sound_score = beauty_scores.get("sound", 7.0)
                para.beauty_word_score = beauty_scores.get("word", 7.0)
                para.beauty_meaning_score = beauty_scores.get("meaning", 7.0)

                success_count += 1

                # 每5段提交一次，避免事务过大
                if (idx + 1) % 5 == 0:
                    db.commit()

            except Exception as e:
                print(f"[Translate Chunk] Error at paragraph {para.paragraph_index}: {e}")
                continue

        db.commit()

        # 全文整体勘误（常识/文化/习俗用语）
        paragraphs = db.query(LiteraryParagraph).filter(
            LiteraryParagraph.translation_id == translation_id
        ).order_by(LiteraryParagraph.paragraph_index).all()
        if paragraphs:
            for batch in _split_paragraphs_into_batches(
                paragraphs,
                lambda p: len(p.step4_finalization or p.translated_text or ""),
            ):
                await _holistic_errata_batch(translation, batch, db)

        paragraphs = db.query(LiteraryParagraph).filter(
            LiteraryParagraph.translation_id == translation_id
        ).order_by(LiteraryParagraph.paragraph_index).all()

        # 更新任务状态
        full_step1 = "\n\n".join([p.step1_translation or "" for p in paragraphs if p.step1_translation])
        full_step2 = "\n\n".join([p.step2_verification or "" for p in paragraphs if p.step2_verification])
        full_step3 = "\n\n".join([p.step3_revision or "" for p in paragraphs if p.step3_revision])
        full_step4 = "\n\n".join([p.step4_finalization or "" for p in paragraphs if p.step4_finalization])

        translation.step1_translation = full_step1
        translation.step2_verification = full_step2
        translation.step3_revision = full_step3
        translation.step4_finalization = full_step4
        translation.final_translation = full_step4
        translation.status = LiteraryTranslationStatus.COMPLETED
        translation.current_step = 4

        from datetime import datetime
        translation.completed_at = datetime.now()

        # 计算平均三美评分
        if paragraphs:
            translation.beauty_sound_score = sum([p.beauty_sound_score or 7.0 for p in paragraphs]) / len(paragraphs)
            translation.beauty_word_score = sum([p.beauty_word_score or 7.0 for p in paragraphs]) / len(paragraphs)
            translation.beauty_meaning_score = sum([p.beauty_meaning_score or 7.0 for p in paragraphs]) / len(paragraphs)

        db.commit()
    finally:
        reset_collab_mode(mode_token)

    # 后台自动提取专业词汇
    background_tasks.add_task(auto_extract_terms_after_finalize, translation_id, db)
    
    return {
        "message": "一键翻译完成",
        "translation_id": translation_id,
        "total_paragraphs": total,
        "success_count": success_count,
        "beauty_scores": {
            "sound": translation.beauty_sound_score,
            "word": translation.beauty_word_score,
            "meaning": translation.beauty_meaning_score
        }
    }


@router.get("/translations/{translation_id}/progress")
async def get_translation_progress(
    translation_id: int,
    db: Session = Depends(get_db)
):
    """
    获取长文本翻译进度
    """
    translation = db.query(LiteraryTranslation).filter(
        LiteraryTranslation.id == translation_id
    ).first()
    
    if not translation:
        raise HTTPException(status_code=404, detail="Translation not found")
    
    paragraphs = db.query(LiteraryParagraph).filter(
        LiteraryParagraph.translation_id == translation_id
    ).all()
    
    total = len(paragraphs)
    completed = sum(1 for p in paragraphs if p.translated_text)
    
    return {
        "translation_id": translation_id,
        "status": translation.status,
        "total_paragraphs": total,
        "completed_paragraphs": completed,
        "progress_percentage": round((completed / total * 100), 1) if total > 0 else 0
    }
